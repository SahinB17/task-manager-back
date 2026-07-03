from datetime import date, datetime, time
from io import BytesIO
import logging
from pathlib import Path
import re
import secrets
import unicodedata
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session

from app import models
from app.auth import get_current_user, hash_password, is_admin_or_manager, require_admin, require_admin_or_manager
from app.database import get_db
from app.utils import create_activity_log
from app.excel_config import (
    GLOBAL_OPERATIONAL_STATUS_MAP,
    GLOBAL_PRIORITY_MAP,
    GLOBAL_STATUS_FROM_OP_STATUS,
    GLOBAL_TASK_COLUMN_ALIASES,
    GLOBAL_TASK_SHEET_ALIASES,
    GLOBAL_TITLE_RECOGNITION_LIST,
)


router = APIRouter(
    prefix="/excel",
    tags=["Excel"]
)

logger = logging.getLogger(__name__)

ALLOWED_TASK_TYPES = {
    "general",
    "one_time",
    "recurring",
    "project",
}
ALLOWED_STATUSES = {"todo", "in_progress", "done"}
ALLOWED_PRIORITIES = {"low", "medium", "high"}
MAX_EXCEL_FILE_SIZE = 10 * 1024 * 1024


EXCEL_OPERATIONAL_STATUS_EXPORT_LABELS = {
    "new": "Yeni",
    "waiting": "Gözləmədə",
    "within_deadline": "İcra müddətinə var",
    "active": "İcradadır",
    "overdue": "Gecikmə",
    "completed": "Tam icra",
    "todo": "Yeni",
    "in_progress": "İcradadır",
    "done": "Tam icra",
}

EXCEL_PRIORITY_EXPORT_LABELS = {
    "low": "Aşağı",
    "medium": "Orta",
    "high": "Yüksək",
}

EXCEL_STAGE_STATUS_EXPORT_LABELS = {
    "not_started": "Başlanmayıb",
    "in_progress": "İcradadır",
    "partial": "Qismən icra",
    "within_deadline": "İcra müddətinə var",
    "waiting": "Gözləmədə",
    "delayed": "Gecikmə",
    "completed": "Tamamlanıb",
}


_AZ_TRANSLIT_MAP = {
    ord("ə"): "a",  # ə -> a
    ord("ı"): "i",  # ı -> i
    ord("ö"): "o",  # ö -> o
    ord("ü"): "u",  # ü -> u
    ord("ş"): "s",  # ş -> s
    ord("ç"): "c",  # ç -> c
    ord("ğ"): "g",  # ğ -> g
    ord("Ə"): "a",  # Ə -> a
    ord("İ"): "i",  # İ -> i
    ord("Ö"): "o",  # Ö -> o
    ord("Ü"): "u",  # Ü -> u
    ord("Ş"): "s",  # Ş -> s
    ord("Ç"): "c",  # Ç -> c
    ord("Ğ"): "g",  # Ğ -> g
}

_AZ_LEGACY_TRANSLIT_MAP = {
    **_AZ_TRANSLIT_MAP,
    ord("ə"): "e",
    ord("Ə"): "e",
}


def normalize_name_whitespace(value):
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def transliterate_az(text):
    """Lowercase a name and transliterate Azerbaijani characters."""
    if not text:
        return ""
    return text.translate(_AZ_TRANSLIT_MAP).lower()


def slugify_az_name(value, translit_map=_AZ_TRANSLIT_MAP):
    normalized_name = normalize_name_whitespace(value)
    transliterated = normalized_name.translate(translit_map).lower()
    ascii_text = unicodedata.normalize("NFKD", transliterated).encode("ascii", "ignore").decode("ascii")
    local_part = re.sub(r"[^a-z0-9]+", ".", ascii_text).strip(".")
    return re.sub(r"\.{2,}", ".", local_part)


def generate_excel_user_email(full_name):
    local_part = slugify_az_name(full_name)
    if not local_part:
        return None
    return f"{local_part}@task-manager.local".lower()


def generate_legacy_excel_user_email(full_name):
    local_part = slugify_az_name(full_name, _AZ_LEGACY_TRANSLIT_MAP)
    if not local_part:
        return None
    return f"{local_part}@task-manager.local".lower()


def get_or_create_user_from_excel(db, full_name):
    normalized_name = normalize_name_whitespace(full_name)
    if not normalized_name:
        return None, False

    generated_email = generate_excel_user_email(normalized_name)
    if not generated_email:
        return None, False

    existing_user = db.query(models.User).filter(
        func.lower(models.User.email) == generated_email.lower()
    ).first()
    if existing_user:
        return existing_user, False

    legacy_email = generate_legacy_excel_user_email(normalized_name)
    if legacy_email and legacy_email != generated_email:
        existing_user = db.query(models.User).filter(
            func.lower(models.User.email) == legacy_email.lower()
        ).first()
        if existing_user:
            return existing_user, False

    user = models.User(
        full_name=normalized_name,
        email=generated_email,
        password_hash=hash_password(secrets.token_urlsafe(32)),
        role_id=3,  # employee
        is_active=False,
        must_change_password=True,
    )
    db.add(user)
    db.flush()
    return user, True


def get_or_create_assignee_users(db, assignee_names):
    """Create or reuse employee users for each unique assignee name.

    Args:
        db: database session
        assignee_names: list of unique non-empty "Icrarci" names
    Returns: (user_map, users_created, users_reused, sample_created_users)
        user_map: dict mapping assignee_name -> models.User
    """
    if not assignee_names:
        return {}, 0, 0, []

    user_map = {}
    users_created = 0
    users_reused = 0
    sample_created_users = []

    for name in assignee_names:
        user, created = get_or_create_user_from_excel(db, name)
        if not user:
            continue

        user_map[name] = user
        if created:
            users_created += 1
            if len(sample_created_users) < 5:
                sample_created_users.append({"name": user.full_name, "email": user.email})
        else:
            users_reused += 1

    return user_map, users_created, users_reused, sample_created_users


def raise_safe_excel_import_error():
    raise HTTPException(
        status_code=500,
        detail="Excel import zamanı xəta baş verdi. Zəhmət olmasa faylı və formatı yoxlayın.",
    )


def normalize_text_key(value):
    normalized_value = unicodedata.normalize("NFKD", value)
    normalized_value = "".join(
        character
        for character in normalized_value
        if not unicodedata.combining(character)
    )
    return normalized_value.casefold().strip()


IMPORT_HEADERS = {
    "task_id",
    "title",
    "description",
    "task_type",
    "status",
    "priority",
    "deadline",
    "client_id",
    "client_name",
    "assigned_to",
    "assigned_user_email",
}

EXPORT_HEADERS = [
    "№",
    "Başlıq",
    "Təsvir",
    "Tapşırıq növü",
    "Status",
    "Prioritet",
    "Son tarix",
    "Müştəri",
    "İcraçı",
    "İcraçı email",
    "Yaradılma tarixi",
    "Yenilənmə tarixi",
    "task_id",
]

TASK_TYPE_LABELS = {
    "general": "Ümumi",
}
STATUS_LABELS = {
    "todo": "Gözləyir",
    "in_progress": "İcradadır",
    "done": "Tamamlandı",
}
PRIORITY_LABELS = {
    "low": "Aşağı",
    "medium": "Orta",
    "high": "Yüksək",
}

TASK_TYPE_IMPORT_ALIASES = {
    normalize_text_key(value): key
    for key, value in TASK_TYPE_LABELS.items()
}
TASK_TYPE_IMPORT_ALIASES.update({
    "umumi": "general",
})
STATUS_IMPORT_ALIASES = {
    normalize_text_key(value): key
    for key, value in STATUS_LABELS.items()
}
STATUS_IMPORT_ALIASES.update({
    "gozleyir": "todo",
    "icradadir": "in_progress",
    "tamamlandi": "done",
})
PRIORITY_IMPORT_ALIASES = {
    normalize_text_key(value): key
    for key, value in PRIORITY_LABELS.items()
}
PRIORITY_IMPORT_ALIASES.update({
    "asagi": "low",
    "orta": "medium",
    "yuksek": "high",
})

HEADER_ALIASES = {
    "başlıq": "title",
    "basliq": "title",
    "təsvir": "description",
    "tesvir": "description",
    "tapşırıq növü": "task_type",
    "tapsiriq novu": "task_type",
    "status": "status",
    "prioritet": "priority",
    "son tarix": "deadline",
    "müştəri": "client_name",
    "musteri": "client_name",
    "icraçı": "assigned_user_email",
    "icraci": "assigned_user_email",
    "icraçı email": "assigned_user_email",
    "icraci email": "assigned_user_email",
}
HEADER_ALIASES = {
    normalize_text_key(key): value
    for key, value in HEADER_ALIASES.items()
}
IGNORED_IMPORT_HEADERS = {
    "№",
    "no",
    "nomre",
    "nömrə",
    "task_id",
    "yaradılma tarixi",
    "yaradilma tarixi",
    "yenilənmə tarixi",
    "yenilenme tarixi",
}
IGNORED_IMPORT_HEADERS = {
    normalize_text_key(header)
    for header in IGNORED_IMPORT_HEADERS
}


def cell_to_str(value):
    if value is None:
        return None

    if isinstance(value, str):
        cleaned_value = value.strip()
        return cleaned_value or None

    return str(value).strip() or None


def preview_cell_value(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value


def read_sheet_records(worksheet):
    rows = worksheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if not header_row:
        return []

    headers = [cell_to_str(value) for value in header_row]
    records = []
    for row_number, row in enumerate(rows, start=2):
        record = {
            header: preview_cell_value(row[index] if index < len(row) else None)
            for index, header in enumerate(headers)
            if header
        }
        if any(value not in (None, "") for value in record.values()):
            record["_row"] = row_number
            records.append(record)
    return records


def parse_yes_no(value):
    normalized_value = normalize_text_key(cell_to_str(value) or "")
    if normalized_value in {"beli", "yes", "true"}:
        return True
    if normalized_value in {"xeyr", "yox", "no", "false"}:
        return False
    return None


def parse_excel_date(value, row_label, col_name, warnings, sheet_name=""):
    """Parse a date from Excel. Handles datetime, date, serial numbers, and text formats."""
    if value is None:
        return None
    # Already a date or datetime
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    # Could be an Excel serial number (integer/float)
    if isinstance(value, (int, float)):
        try:
            from datetime import timedelta
            # Excel serial date: days since 1900-01-01 (with the 1900 leap year bug)
            serial = int(value)
            if serial > 1:
                epoch = date(1899, 12, 30)
                return epoch + timedelta(days=serial)
        except Exception:
            pass
    # Try text formats
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        for fmt in (
            "%m/%d/%Y",
            "%d/%m/%Y",
            "%Y-%m-%d",
            "%d.%m.%Y",
            "%m-%d-%Y",
            "%Y/%m/%d",
        ):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
    # Could not parse
    if sheet_name:
        warnings.append(
            f"[{sheet_name}] Sətir {row_label}: '{col_name}' sütununda tarix formatı tanınmadı: '{value}'"
        )
    else:
        warnings.append(f"Sətir {row_label}: '{col_name}' sütununda tarix formatı tanınmadı: '{value}'")
    return None


# ---------------------------------------------------------------------------
# Flexible Excel column resolution helpers
# ---------------------------------------------------------------------------


def normalize_excel_header(value):
    """Trim spaces, collapse repeated spaces, casefold for comparison.

    Keeps Azerbaijani characters (ə, ı, ö, ü, ş, ç, ğ, İ) intact.
    Handles the Python casefold quirk where "İ" (U+0130) casefolds to
    "i" + combining dot above (U+0307) by normalising to plain "i".
    """
    if value is None:
        return ""
    result = " ".join(str(value).strip().split())
    result = result.casefold()
    # Fix Python casefold for Turkish İ (U+0130):
    # casefold("İ") produces "i̇" — collapse to "i"
    result = result.replace("i̇", "i")
    return result


def normalize_lookup_value(value):
    """Trim, strip, casefold for value comparisons.

    Same normalisation as normalize_excel_header but intended for
    cell *values* rather than column headers.
    """
    if value is None:
        return ""
    result = str(value).strip().casefold()
    result = result.replace("i̇", "i")
    return result


def resolve_excel_column(columns, canonical_field, alias_map):
    """Find the actual column name matching *canonical_field*.

    Args:
        columns: iterable of column name strings (e.g. row dict keys).
        canonical_field: canonical field name (key in *alias_map*).
        alias_map: dict of canonical_field → [primary_header, …aliases].

    Returns:
        The matching column name (str), or None if no match.
    """
    aliases = alias_map.get(canonical_field, [canonical_field])
    if not aliases:
        return None

    # Exact match — try every alias verbatim
    for alias in aliases:
        if alias in columns:
            return alias

    # Normalised match — compare casefolded + space-collapsed forms
    normalized_columns = {}
    for col in columns:
        if col and isinstance(col, str):
            normalized_columns[normalize_excel_header(col)] = col

    for alias in aliases:
        norm_alias = normalize_excel_header(alias)
        if norm_alias in normalized_columns:
            return normalized_columns[norm_alias]

    return None


def build_excel_column_index(columns, alias_map):
    """Pre-resolve every canonical field to its actual column name.

    Returns:
        dict: canonical_field → actual_column_name (only resolved fields).
    """
    index = {}
    for canonical_field in alias_map:
        col = resolve_excel_column(columns, canonical_field, alias_map)
        if col is not None:
            index[canonical_field] = col
    return index


def collect_unknown_excel_columns(columns, alias_map, col_index):
    """Identify column headers that did not match any canonical field alias.

    Args:
        columns: iterable of actual column header strings from the sheet.
        alias_map: canonical_field → [primary_header, ...aliases] mapping.
        col_index: pre-resolved index from build_excel_column_index().

    Returns:
        Sorted list of unknown column header strings.
    """
    if not columns:
        return []

    # Collect all known normalized aliases
    known_normalized = set()
    for canonical_field, aliases in alias_map.items():
        for alias in aliases:
            known_normalized.add(normalize_excel_header(alias))
        if col_index and canonical_field in col_index:
            actual_col = col_index[canonical_field]
            if actual_col:
                known_normalized.add(normalize_excel_header(actual_col))

    unknown = []
    for col in columns:
        if col is None or not isinstance(col, str) or not col.strip():
            continue
        if col.startswith("_"):
            continue
        norm = normalize_excel_header(col)
        if norm not in known_normalized:
            unknown.append(col)

    return sorted(unknown)


def get_excel_value(row, canonical_field, alias_map,
                  col_index=None, default=None):
    """Get a cell value from a row dict using column aliases.

    Args:
        row: dict of column_name → cell_value.
        canonical_field: canonical field name to look up.
        alias_map: dict of canonical_field → [primary_header, …aliases].
        col_index: optional pre-resolved index from build_excel_column_index().
        default: returned when the column cannot be resolved.

    Returns:
        The cell value, or *default*.
    """
    if col_index is not None:
        col = col_index.get(canonical_field)
    else:
        col = resolve_excel_column(row.keys(), canonical_field, alias_map)

    if col is None:
        return default
    return row.get(col)


def normalize_excel_map_value(raw_value, value_alias_map):
    """Map a raw Excel value through value aliases to a canonical AZ label.

    Args:
        raw_value: raw cell value (str or None).
        value_alias_map: dict of normalised_alias → canonical_azeri_label.

    Returns:
        The canonical Azerbaijani label if an alias matches, otherwise
        the original *raw_value* unchanged (pass-through for back-compat).
    """
    if raw_value is None:
        return None
    norm = normalize_lookup_value(raw_value)
    return value_alias_map.get(norm, raw_value)


# ---------------------------------------------------------------------------
# Phase 2: Universal Excel sheet detection helpers
# ---------------------------------------------------------------------------


def normalize_sheet_name(value):
    """Trim, collapse spaces, casefold for sheet name comparison.

    Azerbaijani-safe: handles the Python casefold quirk where
    "İ" (U+0130) casefolds to "i" + combining dot above (U+0307).
    """
    if value is None:
        return ""
    result = " ".join(str(value).strip().split())
    result = result.casefold()
    result = result.replace("i̇", "i")
    return result


def normalize_header(value):
    header = normalize_text_key(cell_to_str(value) or "")
    return HEADER_ALIASES.get(header, header)


def format_excel_datetime(value):
    if not value:
        return None

    return value.strftime("%d.%m.%Y %H:%M")


def format_excel_export_date(value):
    if isinstance(value, datetime):
        return value.date()
    return value


def format_excel_yes_no(value):
    if value is None:
        return None
    return "Bəli" if value else "Xeyr"


def get_excel_status_fill(value):
    colors = {
        "Yeni": "DBEAFE",
        "Gözləmədə": "F3F4F6",
        "İcra müddətinə var": "FEF3C7",
        "İcradadır": "DCFCE7",
        "Gecikmə": "FEE2E2",
        "Tam icra": "D1FAE5",
        "Başlanmayıb": "F3F4F6",
        "Qismən icra": "FEF3C7",
        "Tamamlanıb": "D1FAE5",
    }
    color = colors.get(value)
    return PatternFill(fill_type="solid", fgColor=color) if color else None


def apply_excel_export_styles(workbook):
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin", color="D9E2EC"),
        right=Side(style="thin", color="D9E2EC"),
        top=Side(style="thin", color="D9E2EC"),
        bottom=Side(style="thin", color="D9E2EC"),
    )
    zebra_fill = PatternFill(fill_type="solid", fgColor="F8FAFC")

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        if worksheet.max_row and worksheet.max_column:
            worksheet.auto_filter.ref = worksheet.dimensions

        status_column = None
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
            if cell.value == "Status":
                status_column = cell.column
        worksheet.row_dimensions[1].height = 28

        for row in worksheet.iter_rows(min_row=2):
            max_text_length = 0
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = thin_border
                if cell.row % 2 == 0:
                    cell.fill = zebra_fill

                value_length = len(str(cell.value)) if cell.value is not None else 0
                max_text_length = max(max_text_length, value_length)

                if status_column and cell.column == status_column:
                    status_fill = get_excel_status_fill(cell.value)
                    if status_fill:
                        cell.fill = status_fill
                        cell.font = Font(bold=True, color="111827")

            if max_text_length > 120:
                worksheet.row_dimensions[row[0].row].height = 64
            elif max_text_length > 70:
                worksheet.row_dimensions[row[0].row].height = 46
            elif max_text_length > 40:
                worksheet.row_dimensions[row[0].row].height = 32
            else:
                worksheet.row_dimensions[row[0].row].height = 22

        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 3, 12), 45)


def parse_uuid(value, field_name, row_number, errors):
    if not value:
        return None

    try:
        return UUID(str(value))
    except ValueError:
        errors.append({
            "row": row_number,
            "field": field_name,
            "message": f"{field_name} düzgün UUID olmalıdır"
        })
        return None


def find_existing_task(db, task_id, title, client_id, deadline):
    if task_id:
        return db.query(models.Task).filter(models.Task.id == task_id).first()

    query = db.query(models.Task).filter(
        models.Task.title == title,
        models.Task.client_id == client_id,
    )

    if deadline is None:
        query = query.filter(models.Task.deadline.is_(None))
    else:
        query = query.filter(models.Task.deadline == deadline)

    return query.first()


def parse_deadline(value, row_number, errors):
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return datetime.combine(value, time.min)

    value_text = cell_to_str(value)
    if not value_text:
        return None

    normalized_value = value_text.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(normalized_value)
    except ValueError:
        pass

    for date_format in (
        "%Y-%m-%d",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%d.%m.%Y",
        "%d.%m.%Y %H:%M",
        "%d.%m.%Y %H:%M:%S",
        "%d/%m/%Y",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y %H:%M:%S",
    ):
        try:
            return datetime.strptime(value_text, date_format)
        except ValueError:
            continue

    errors.append({
        "row": row_number,
        "field": "deadline",
        "message": "deadline düzgün tarix və ya tarix-saat formatında olmalıdır"
    })
    return None


def validate_choice(value, allowed_values, default_value, field_name, row_number, errors):
    normalized_value = (cell_to_str(value) or default_value).lower()

    if normalized_value not in allowed_values:
        errors.append({
            "row": row_number,
            "field": field_name,
            "message": f"{field_name} bu dəyərlərdən biri olmalıdır: {', '.join(sorted(allowed_values))}"
        })

    return normalized_value


def validate_task_type(value, row_number, errors):
    normalized_value = cell_to_str(value)
    if not normalized_value:
        return "general"

    value_key = normalize_text_key(normalized_value)
    value_key = TASK_TYPE_IMPORT_ALIASES.get(value_key, value_key)

    return validate_choice(value_key, ALLOWED_TASK_TYPES, "general", "task_type", row_number, errors)


def validate_status(value, row_number, errors):
    normalized_value = cell_to_str(value)
    if not normalized_value:
        return "todo"

    value_key = normalize_text_key(normalized_value)
    value_key = STATUS_IMPORT_ALIASES.get(value_key, value_key)

    return validate_choice(value_key, ALLOWED_STATUSES, "todo", "status", row_number, errors)


def validate_priority(value, row_number, errors):
    normalized_value = cell_to_str(value)
    if not normalized_value:
        return "medium"

    value_key = normalize_text_key(normalized_value)
    value_key = PRIORITY_IMPORT_ALIASES.get(value_key, value_key)

    return validate_choice(value_key, ALLOWED_PRIORITIES, "medium", "priority", row_number, errors)


def validate_excel_file(file):
    safe_name = Path(file.filename or "").name
    extension = Path(safe_name).suffix.lower()

    if extension != ".xlsx":
        raise HTTPException(status_code=400, detail="Yalnız .xlsx Excel faylları dəstəklənir")

    file.file.seek(0, 2)
    file_size = file.file.tell()
    file.file.seek(0)

    if file_size > MAX_EXCEL_FILE_SIZE:
        raise HTTPException(
            status_code=400,
            detail="Fayl ölçüsü çox böyükdür. Maksimum icazəli ölçü 10 MB-dır"
        )


def validate_excel_file_size(file):
    """Check that an uploaded Excel file does not exceed the size limit."""
    file.file.seek(0, 2)
    file_size = file.file.tell()
    file.file.seek(0)

    if file_size > MAX_EXCEL_FILE_SIZE:
        raise HTTPException(
            status_code=400,
            detail="Fayl ölçüsü çox böyükdür. Maksimum icazəli ölçü 10 MB-dır"
        )


def find_client(db, client_id, client_name):
    if client_id:
        return db.query(models.Client).filter(models.Client.id == client_id).first()

    if client_name:
        return db.query(models.Client).filter(models.Client.company_name == client_name).first()

    return None


def find_assigned_user(db, assigned_to, assigned_user_email):
    if assigned_to:
        return db.query(models.User).filter(models.User.id == assigned_to).first()

    if assigned_user_email:
        return db.query(models.User).filter(models.User.email == assigned_user_email).first()

    return None


@router.get("/tasks/export")
def export_tasks_to_excel(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    if is_admin_or_manager(current_user):
        tasks = db.query(models.Task).order_by(models.Task.created_at).all()
        projects = db.query(models.Project).order_by(models.Project.name).all()
        stages = db.query(models.ProjectStage).join(
            models.Project,
            models.ProjectStage.project_id == models.Project.id,
        ).order_by(
            models.Project.name,
            models.ProjectStage.sort_order,
            models.ProjectStage.created_at,
        ).all()
    else:
        tasks = db.query(models.Task).filter(
            models.Task.assigned_to == current_user.id
        ).order_by(models.Task.created_at).all()
        project_ids = {task.project_id for task in tasks if task.project_id}
        projects = (
            db.query(models.Project)
            .filter(models.Project.id.in_(project_ids))
            .order_by(models.Project.name)
            .all()
            if project_ids
            else []
        )
        stages = (
            db.query(models.ProjectStage)
            .join(models.Project, models.ProjectStage.project_id == models.Project.id)
            .filter(models.ProjectStage.project_id.in_(project_ids))
            .order_by(
                models.Project.name,
                models.ProjectStage.sort_order,
                models.ProjectStage.created_at,
            )
            .all()
            if project_ids
            else []
        )

    workbook = Workbook()
    master_sheet = workbook.active
    master_sheet.title = "Tapşırıqlar"
    project_sheet = workbook.create_sheet("Qruplar")
    stage_sheet = workbook.create_sheet("Mərhələlər")

    custom_field_columns = []
    for task in tasks:
        if not isinstance(task.custom_fields, dict):
            continue

        for key in task.custom_fields:
            if key and key not in custom_field_columns:
                custom_field_columns.append(key)

    master_sheet.append([
        "Xarici ID",
        "Tapşırıq",
        "Təsvir",
        "İcraçı",
        "Prioritet",
        "Status",
        "Başlama tarixi",
        "Son tarix",
        "Kateqoriya",
        "Kanban",
        *custom_field_columns,
    ])
    project_sheet.append([
        "Layihə_İD",
        "Layihənin adı",
        "Kateqoriya",
        "Sahibi",
        "Yanaşma",
    ])
    stage_sheet.append([
        "Layihə",
        "Mərhələ",
        "Çəki %",
        "Plan Başlama",
        "Plan Bitmə",
        "Fakt Bitmə",
        "Status",
        "Ləngimə səbəbi",
        "Mərhələnin icra faizi",
    ])

    for task in tasks:
        task_custom_fields = task.custom_fields if isinstance(task.custom_fields, dict) else {}
        master_sheet.append([
            task.external_id or task.sxa_task_id or str(task.id),
            task.title,
            task.description,
            task.assigned_user.full_name if task.assigned_user else None,
            EXCEL_PRIORITY_EXPORT_LABELS.get(task.priority, task.priority),
            EXCEL_OPERATIONAL_STATUS_EXPORT_LABELS.get(
                task.operational_status or task.status,
                task.operational_status or task.status,
            ),
            format_excel_export_date(task.start_date or task.sxa_start_date),
            format_excel_export_date(task.current_deadline or task.deadline),
            task.category or task.direction_program or (task.project.direction_program if task.project else None),
            format_excel_yes_no(task.kanban_enabled),
            *[task_custom_fields.get(column) for column in custom_field_columns],
        ])

    for project in projects:
        project_sheet.append([
            project.code,
            project.name,
            project.direction_program,
            project.owner,
            project.approach,
        ])

    for stage in stages:
        stage_sheet.append([
            stage.project.name if stage.project else None,
            stage.name,
            stage.weight_percent,
            format_excel_export_date(stage.planned_start),
            format_excel_export_date(stage.planned_end),
            format_excel_export_date(stage.actual_end),
            EXCEL_STAGE_STATUS_EXPORT_LABELS.get(stage.status, stage.status),
            stage.delay_reason,
            stage.completion_percent,
        ])

    apply_excel_export_styles(workbook)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"task-manager-export-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.post("/tasks/import")
def import_tasks_from_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin)
):
    validate_excel_file(file)

    try:
        workbook = load_workbook(file.file, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Excel faylı oxuna bilmədi")

    worksheet = workbook.active
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)

    if not header_row:
        raise HTTPException(status_code=400, detail="Excel faylında başlıq sətri olmalıdır")

    headers = [normalize_header(header) for header in header_row]
    unknown_headers = [
        header
        for header in headers
        if header and header not in IMPORT_HEADERS and header not in IGNORED_IMPORT_HEADERS
    ]
    missing_headers = [header for header in ("title",) if header not in headers]

    if missing_headers:
        raise HTTPException(
            status_code=400,
            detail={
                "message": "Excel başlıqları yanlışdır",
                "missing_headers": missing_headers,
                "unknown_headers": unknown_headers,
                "allowed_headers": sorted(IMPORT_HEADERS)
            }
        )

    errors = []
    task_payloads = []
    skipped_empty_rows = 0

    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        row_data = {
            headers[index]: row[index]
            for index in range(min(len(headers), len(row)))
            if headers[index]
        }

        if not any(value not in (None, "") for value in row_data.values()):
            skipped_empty_rows += 1
            continue

        title = cell_to_str(row_data.get("title"))
        if not title:
            errors.append({
                "row": row_number,
                "field": "title",
                "message": "title tələb olunur"
            })

        task_id = parse_uuid(cell_to_str(row_data.get("task_id")), "task_id", row_number, errors)
        client_id = parse_uuid(cell_to_str(row_data.get("client_id")), "client_id", row_number, errors)
        assigned_to = parse_uuid(cell_to_str(row_data.get("assigned_to")), "assigned_to", row_number, errors)
        client_name = cell_to_str(row_data.get("client_name"))
        assigned_user_email = cell_to_str(row_data.get("assigned_user_email"))

        client = find_client(db, client_id, client_name)
        if not client:
            errors.append({
                "row": row_number,
                "field": "client_id/client_name",
                "message": "Mövcud müştəri tələb olunur"
            })

        assigned_user = find_assigned_user(db, assigned_to, assigned_user_email)
        if (assigned_to or assigned_user_email) and not assigned_user:
            errors.append({
                "row": row_number,
                "field": "assigned_to/assigned_user_email",
                "message": "Təyin edilən istifadəçi tapılmadı"
            })

        task_type = validate_task_type(row_data.get("task_type"), row_number, errors)
        status = validate_status(row_data.get("status"), row_number, errors)
        priority = validate_priority(row_data.get("priority"), row_number, errors)
        deadline = parse_deadline(row_data.get("deadline"), row_number, errors)

        existing_task = None
        if title and client:
            existing_task = find_existing_task(db, task_id, title, client.id, deadline)

        if task_id and not existing_task:
            errors.append({
                "row": row_number,
                "field": "task_id",
                "message": "Mövcud task tapılmadı"
            })

        if title and client:
            task_payloads.append({
                "existing_task": existing_task,
                "title": title,
                "description": cell_to_str(row_data.get("description")),
                "task_type": task_type,
                "status": status,
                "priority": priority,
                "deadline": deadline,
                "client_id": client.id,
                "assigned_to": assigned_user.id if assigned_user else None,
            })

    if errors:
        raise HTTPException(
            status_code=400,
            detail={
                "message": "Excel import validation xətası",
                "errors": errors,
                "imported_count": 0,
                "skipped_empty_rows": skipped_empty_rows
            }
        )

    created_tasks = []
    updated_tasks = []
    for payload in task_payloads:
        task = payload["existing_task"]

        if task:
            old_data = {
                "title": task.title,
                "description": task.description,
                "task_type": task.task_type,
                "status": task.status,
                "priority": task.priority,
                "deadline": str(task.deadline) if task.deadline else None,
                "client_id": str(task.client_id) if task.client_id else None,
                "assigned_to": str(task.assigned_to) if task.assigned_to else None,
            }

            task.title = payload["title"]
            task.description = payload["description"]
            task.task_type = payload["task_type"]
            task.status = payload["status"]
            task.priority = payload["priority"]
            task.deadline = payload["deadline"]
            task.client_id = payload["client_id"]
            task.assigned_to = payload["assigned_to"]

            create_activity_log(
                db=db,
                user_id=current_user.id,
                action="TASK_IMPORT_UPDATED",
                entity_type="task",
                entity_id=task.id,
                old_data=old_data,
                new_data={
                    "title": task.title,
                    "description": task.description,
                    "task_type": task.task_type,
                    "status": task.status,
                    "priority": task.priority,
                    "deadline": str(task.deadline) if task.deadline else None,
                    "client_id": str(task.client_id) if task.client_id else None,
                    "assigned_to": str(task.assigned_to) if task.assigned_to else None,
                }
            )
            updated_tasks.append(task)
        else:
            task = models.Task(
                title=payload["title"],
                description=payload["description"],
                task_type=payload["task_type"],
                status=payload["status"],
                priority=payload["priority"],
                deadline=payload["deadline"],
                client_id=payload["client_id"],
                assigned_to=payload["assigned_to"],
                created_by=current_user.id
            )

            db.add(task)
            db.flush()

            create_activity_log(
                db=db,
                user_id=current_user.id,
                action="TASK_IMPORTED",
                entity_type="task",
                entity_id=task.id,
                old_data=None,
                new_data={
                    "title": task.title,
                    "description": task.description,
                    "task_type": task.task_type,
                    "status": task.status,
                    "priority": task.priority,
                    "deadline": str(task.deadline) if task.deadline else None,
                    "client_id": str(task.client_id) if task.client_id else None,
                    "assigned_to": str(task.assigned_to) if task.assigned_to else None,
                    "created_by": str(task.created_by) if task.created_by else None
                }
            )
            created_tasks.append(task)

    db.commit()

    return {
        "message": "Tasklar import edildi",
        "imported_count": len(created_tasks),
        "updated_count": len(updated_tasks),
        "skipped_empty_rows": skipped_empty_rows
    }


# =============================================================================
# Global Excel Import — task-first, no projects/stages required
# =============================================================================

# Pre-computed set of normalised title column headers for fallback detection
_GLOBAL_TITLE_RECOGNITION_HEADERS = frozenset(
    normalize_excel_header(h) for h in GLOBAL_TITLE_RECOGNITION_LIST
)


def find_global_task_sheet(workbook):
    """Find the task sheet in a global-import workbook.

    Priority:
      1. If only 1 sheet → use it.
      2. If multiple sheets → prefer names matching GLOBAL_TASK_SHEET_ALIASES.
      3. If none match → first sheet with a recognizable title column header.

    Returns:
        openpyxl Worksheet

    Raises:
        HTTPException 400 when no suitable task sheet can be found.
    """
    if len(workbook.sheetnames) == 1:
        return workbook[workbook.sheetnames[0]]

    # Prefer known sheet names
    normalized_aliases = {normalize_sheet_name(a) for a in GLOBAL_TASK_SHEET_ALIASES}
    for sheet_name in workbook.sheetnames:
        if normalize_sheet_name(sheet_name) in normalized_aliases:
            return workbook[sheet_name]

    # Fallback: first sheet with a recognizable title column
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        header_row = next(ws.iter_rows(values_only=True), None)
        if not header_row:
            continue
        for value in header_row:
            if value is not None and normalize_excel_header(str(value)) in _GLOBAL_TITLE_RECOGNITION_HEADERS:
                return ws

    raise HTTPException(
        status_code=400,
        detail={
            "message": (
                "Tapşırıq vərəqi tapılmadı. "
                "Faylda tapşırıq başlığı olan sütun olmalıdır "
                "(məs: Tapşırıq, Başlıq, Title, Task)."
            ),
            "required_missing": ["title"],
        },
    )


def _build_custom_fields_from_row(row, col_index, alias_map):
    """Collect unknown column values from a row into a custom_fields dict.

    Known columns (those resolved in *col_index*) are excluded.
    Headers that are empty, start with "_", or look like "Unnamed: N" are skipped.
    Empty / None / NaN values are skipped.

    Returns:
        dict of {original_header: cell_value}, or empty dict.
    """
    custom_fields = {}

    # Build set of resolved actual column names for exclusion
    resolved_names = set(col_index.values())

    for header, value in row.items():
        if header is None or header == "_row":
            continue
        header_str = str(header).strip()
        if not header_str:
            continue
        if header_str.startswith("_"):
            continue
        # Skip "Unnamed: N" pattern (pandas-style default headers)
        import re
        if re.match(r"^Unnamed:\s*\d+$", header_str, re.IGNORECASE):
            continue
        # Skip known mapped columns
        if header_str in resolved_names:
            continue
        # Skip if also mapped by normalized match
        normalized = normalize_excel_header(header_str)
        is_known = False
        for known_col in resolved_names:
            if normalize_excel_header(known_col) == normalized:
                is_known = True
                break
        if is_known:
            continue

        # Value handling
        if value is None:
            continue
        # Handle NaN float
        if isinstance(value, float):
            import math
            if math.isnan(value):
                continue
            # Simple float — keep as is or convert to int if whole number
            if value == int(value):
                custom_fields[header_str] = int(value)
            else:
                custom_fields[header_str] = value
            continue
        # Dates → ISO string
        if isinstance(value, (date, datetime)):
            custom_fields[header_str] = (
                value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
            )
            continue
        # Boolean → "Bəli" / "Xeyr"
        if isinstance(value, bool):
            custom_fields[header_str] = "Bəli" if value else "Xeyr"
            continue
        # String — strip, skip if empty
        if isinstance(value, str):
            stripped = value.strip()
            if not stripped or stripped.lower() == "nan":
                continue
            custom_fields[header_str] = stripped
            continue
        # Other types — convert to string
        s = str(value).strip()
        if s and s.lower() != "nan":
            custom_fields[header_str] = s

    return custom_fields


def _auto_create_category_option(db, raw_name):
    """Create a managed category option if it doesn't already exist.

    Called during global Excel import to ensure imported categories
    appear in the managed options list. Idempotent — skips existing.
    """
    if not raw_name or not raw_name.strip():
        return
    from app.models import CategoryOption
    from sqlalchemy import func as _fn

    trimmed = raw_name.strip()
    normalized = " ".join(trimmed.split()).lower()
    existing = (
        db.query(CategoryOption)
        .filter(_fn.lower(CategoryOption.normalized_name) == normalized)
        .first()
    )
    if existing:
        if not existing.is_active:
            existing.is_active = True
            existing.name = trimmed
            existing.normalized_name = normalized
        return
    option = CategoryOption(name=trimmed, normalized_name=normalized, is_active=True, sort_order=0)
    db.add(option)


def _has_row_changes(existing_task, title, description, category_raw, priority,
                     op_status, derived_status, kanban_enabled, deadline_dt,
                     start_date_date, external_id_raw, assigned_user_id,
                     row_custom_fields):
    """Compare extracted Excel row values against an existing task.

    Returns True if at least one field would actually change the task.
    Empty/none Excel values are treated as "no intent to change" —
    they do not count as a change.
    """
    if title and title != existing_task.title:
        return True
    if description and description != (existing_task.description or ""):
        return True
    if category_raw and category_raw != (existing_task.category or existing_task.direction_program or ""):
        return True
    if priority and priority != (existing_task.priority or "medium"):
        return True
    if op_status and op_status != (existing_task.operational_status or "new"):
        return True
    if kanban_enabled is not None and kanban_enabled != (existing_task.kanban_enabled if existing_task.kanban_enabled is not None else True):
        return True
    if deadline_dt is not None:
        existing_dl = existing_task.deadline
        if existing_dl is None or deadline_dt != existing_dl:
            return True
    if start_date_date is not None:
        existing_sd = existing_task.start_date
        if existing_sd is None or start_date_date != existing_sd:
            return True
    if external_id_raw and external_id_raw.strip() != (existing_task.external_id or ""):
        return True
    if assigned_user_id is not None:
        existing_au = str(existing_task.assigned_to) if existing_task.assigned_to else None
        if str(assigned_user_id) != existing_au:
            return True
    # custom_fields: merge Excel into existing, compare result
    if row_custom_fields:
        merged = dict(existing_task.custom_fields or {})
        merged.update(row_custom_fields)
        if merged != (existing_task.custom_fields or {}):
            return True

    return False


def _match_task_for_import(db, external_id, title, assignee_name, deadline_date):
    """Find an existing task matching the given row data.

    Returns:
        dict with keys: action, matched_task, match_reason, warnings
        action is one of: "create", "update", "ambiguous"
    """
    warnings = []

    # Rule 1: external_id match
    if external_id and external_id.strip():
        matches = (
            db.query(models.Task)
            .filter(func.lower(models.Task.external_id) == func.lower(external_id.strip()))
            .all()
        )
        if len(matches) == 1:
            return {"action": "update", "matched_task": matches[0], "match_reason": "external_id", "warnings": warnings}
        if len(matches) > 1:
            warnings.append(
                f"'{external_id}' Xarici ID-si ilə {len(matches)} task tapıldı — birmənalı seçim mümkün deyil."
            )
            return {"action": "ambiguous", "matched_task": None, "match_reason": "external_id_ambiguous", "warnings": warnings}
        # No match → create
        return {"action": "create", "matched_task": None, "match_reason": "external_id", "warnings": warnings}

    # Rule 2: strict title + assignee + deadline fallback
    if title and assignee_name and deadline_date is not None:
        # Resolve assignee to user_id for matching
        assignee_user = (
            db.query(models.User)
            .filter(func.lower(models.User.full_name) == func.lower(assignee_name.strip()))
            .first()
        )
        assignee_id = assignee_user.id if assignee_user else None

        query = db.query(models.Task).filter(
            models.Task.title == title,
            models.Task.deadline == datetime.combine(deadline_date, time.min),
        )
        if assignee_id:
            query = query.filter(models.Task.assigned_to == assignee_id)
        else:
            query = query.filter(models.Task.assigned_to.is_(None))

        matches = query.all()
        if len(matches) == 1:
            return {"action": "update", "matched_task": matches[0], "match_reason": "title_assignee_deadline", "warnings": warnings}
        if len(matches) > 1:
            warnings.append(
                f"'{title}' başlığı, '{assignee_name}' icraçısı ilə {len(matches)} task tapıldı — birmənalı seçim mümkün deyil."
            )
            return {"action": "ambiguous", "matched_task": None, "match_reason": "title_assignee_deadline_ambiguous", "warnings": warnings}
        # No match → create
        return {"action": "create", "matched_task": None, "match_reason": "title_assignee_deadline", "warnings": warnings}

    # No safe match possible → create
    return {"action": "create", "matched_task": None, "match_reason": "no_match_key", "warnings": warnings}


def _apply_preview_matching(preview_data, db):
    """Add action/match fields to each task in the preview result.

    Modifies *preview_data* in place — adds action, matched_task_id,
    match_reason to each task, and adds summary counts.
    """
    create_count = 0
    update_count = 0
    no_change_count = 0
    ambiguous_count = 0
    skipped_count = 0

    for t in preview_data.get("tasks", []):
        external_id = t.get("external_id")
        title = t.get("title")
        assignee_name = t.get("assignee")
        deadline_raw = t.get("deadline")
        deadline_date = None
        if deadline_raw:
            if isinstance(deadline_raw, (date, datetime)):
                deadline_date = deadline_raw if isinstance(deadline_raw, date) else deadline_raw.date()
            elif isinstance(deadline_raw, str):
                try:
                    deadline_date = date.fromisoformat(deadline_raw)
                except ValueError:
                    pass

        match = _match_task_for_import(db, external_id, title, assignee_name, deadline_date)

        # Detect no-change: if matched, compare extracted values against existing task
        if match["action"] == "update" and match["matched_task"]:
            existing = match["matched_task"]
            # Map raw priority / op_status from preview dict to internal values
            priority_raw = t.get("priority")
            pr_norm = normalize_lookup_value(priority_raw) if priority_raw else ""
            mapped_priority = GLOBAL_PRIORITY_MAP.get(pr_norm, "medium")
            op_raw = t.get("operational_status")
            op_norm = normalize_lookup_value(op_raw) if op_raw else ""
            mapped_op = GLOBAL_OPERATIONAL_STATUS_MAP.get(op_norm, "new")
            mapped_status = GLOBAL_STATUS_FROM_OP_STATUS.get(mapped_op, "todo")
            # Parse start_date from preview
            sd_raw = t.get("start_date")
            start_date_date_preview = None
            if sd_raw:
                if isinstance(sd_raw, date):
                    start_date_date_preview = sd_raw
                elif isinstance(sd_raw, str):
                    try:
                        start_date_date_preview = date.fromisoformat(sd_raw)
                    except ValueError:
                        pass
            # Parse kanban
            kanb_raw = t.get("kanban")
            kanb_val = parse_yes_no(kanb_raw)

            changed = _has_row_changes(
                existing, title, t.get("description"),
                t.get("category"), mapped_priority,
                mapped_op, mapped_status, kanb_val,
                datetime.combine(deadline_date, time.min) if deadline_date else None,
                start_date_date_preview, external_id,
                None,  # assignee not resolved in preview
                t.get("custom_fields") or {},
            )
            if not changed:
                match["action"] = "no_change"

        t["action"] = match["action"]
        t["matched_task_id"] = str(match["matched_task"].id) if match["matched_task"] else None
        t["match_reason"] = match["match_reason"]
        if match["warnings"]:
            preview_data.setdefault("warnings", []).extend(match["warnings"])

        if match["action"] == "create":
            create_count += 1
        elif match["action"] == "update":
            update_count += 1
        elif match["action"] == "no_change":
            no_change_count += 1
        elif match["action"] == "ambiguous":
            ambiguous_count += 1
        else:
            skipped_count += 1

    preview_data["create_count"] = create_count
    preview_data["update_count"] = update_count
    preview_data["no_change_count"] = no_change_count
    preview_data["ambiguous_count"] = ambiguous_count
    preview_data["skipped_count"] = skipped_count


def parse_global_preview(file_bytes, filename=""):
    """Parse a global Excel workbook for preview — no database writes.

    Returns:
        dict with keys: mode, tasks_count, tasks, warnings,
                        unknown_columns, required_missing
    """
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Excel faylı oxuna bilmədi") from exc

    try:
        task_ws = find_global_task_sheet(workbook)
        sheet_name = task_ws.title
        rows = read_sheet_records(task_ws)
    finally:
        workbook.close()

    warnings = []
    unknown_columns = []
    required_missing = []

    if not rows:
        warnings.append("Tapşırıq sətirləri tapılmadı.")
        return {
            "mode": "global",
            "tasks_count": 0,
            "tasks": [],
            "warnings": warnings,
            "unknown_columns": unknown_columns,
            "required_missing": required_missing,
        }

    col_index = build_excel_column_index(rows[0].keys(), GLOBAL_TASK_COLUMN_ALIASES)

    # Title column is required
    if "title" not in col_index:
        required_missing.append("title")
        return {
            "mode": "global",
            "tasks_count": 0,
            "tasks": [],
            "warnings": [
                "Başlıq (Tapşırıq / Title / Task) sütunu tapılmadı. "
                "Bu sütun tələb olunur."
            ],
            "unknown_columns": unknown_columns,
            "required_missing": required_missing,
        }

    # Detect unknown columns
    unknown_columns = collect_unknown_excel_columns(
        list(rows[0].keys()), GLOBAL_TASK_COLUMN_ALIASES, col_index
    )
    if unknown_columns:
        warnings.append(
            f"'{sheet_name}' vərəqində əlavə sahə kimi saxlanacaq sütun(lar): "
            f"{', '.join(unknown_columns)}. "
            f"Bu sütunlar əsas sahə deyil, əlavə sahə kimi saxlanacaq."
        )

    tasks = []
    skipped_empty_title = 0
    custom_field_columns = set()

    for row in rows:
        title = cell_to_str(
            get_excel_value(row, "title", GLOBAL_TASK_COLUMN_ALIASES, col_index)
        )
        if not title:
            skipped_empty_title += 1
            warnings.append(
                f"Sətir {row['_row']}: Başlıq boşdur — bu sətir keçiriləcək."
            )
            continue

        # Collect unknown columns into custom_fields
        row_custom_fields = _build_custom_fields_from_row(
            row, col_index, GLOBAL_TASK_COLUMN_ALIASES
        )
        for key in row_custom_fields:
            custom_field_columns.add(key)

        task_record = {
            "_row": row["_row"],
            "title": title,
            "description": cell_to_str(
                get_excel_value(row, "description", GLOBAL_TASK_COLUMN_ALIASES, col_index)
            ),
            "assignee": cell_to_str(
                get_excel_value(row, "assignee", GLOBAL_TASK_COLUMN_ALIASES, col_index)
            ),
            "priority": cell_to_str(
                get_excel_value(row, "priority", GLOBAL_TASK_COLUMN_ALIASES, col_index)
            ),
            "operational_status": cell_to_str(
                get_excel_value(row, "operational_status", GLOBAL_TASK_COLUMN_ALIASES, col_index)
            ),
            "deadline": get_excel_value(
                row, "deadline", GLOBAL_TASK_COLUMN_ALIASES, col_index
            ),
            "start_date": get_excel_value(
                row, "start_date", GLOBAL_TASK_COLUMN_ALIASES, col_index
            ),
            "category": cell_to_str(
                get_excel_value(row, "category", GLOBAL_TASK_COLUMN_ALIASES, col_index)
            ),
            "kanban": cell_to_str(
                get_excel_value(row, "kanban", GLOBAL_TASK_COLUMN_ALIASES, col_index)
            ),
            "external_id": cell_to_str(
                get_excel_value(row, "external_id", GLOBAL_TASK_COLUMN_ALIASES, col_index)
            ),
            "custom_fields": row_custom_fields,
        }
        tasks.append(task_record)

    if skipped_empty_title:
        warnings.append(
            f"{skipped_empty_title} sətir başlıq olmadığı üçün keçirildi."
        )

    return {
        "mode": "global",
        "tasks_count": len(tasks),
        "tasks": tasks,
        "warnings": warnings,
        "unknown_columns": unknown_columns,
        "custom_field_columns": sorted(custom_field_columns),
        "required_missing": required_missing,
    }


def execute_global_import(file_bytes, db, current_user):
    """Parse and import tasks from a global Excel workbook.

    Row-level errors add warnings and skip the row — they never abort
    the entire import.

    Returns:
        dict with keys: mode, created_tasks, updated_tasks, skipped_rows,
                        tasks, warnings, unknown_columns
    """
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Excel faylı oxuna bilmədi") from exc

    try:
        task_ws = find_global_task_sheet(workbook)
        sheet_name = task_ws.title
        rows = read_sheet_records(task_ws)
    finally:
        workbook.close()

    warnings = []
    unknown_columns = []

    if not rows:
        return {
            "mode": "global",
            "created_tasks": 0,
            "updated_tasks": 0,
            "no_change_count": 0,
            "skipped_rows": 0,
            "ambiguous_rows": 0,
            "tasks": [],
            "warnings": ["Tapşırıq sətirləri tapılmadı."],
            "unknown_columns": [],
            "custom_field_columns": [],
        }

    col_index = build_excel_column_index(rows[0].keys(), GLOBAL_TASK_COLUMN_ALIASES)

    # Title column is required
    if "title" not in col_index:
        raise HTTPException(
            status_code=400,
            detail={
                "message": (
                    "Başlıq (Tapşırıq / Title / Task) sütunu tapılmadı. "
                    "Bu sütun tələb olunur."
                ),
                "required_missing": ["title"],
            },
        )

    unknown_columns = collect_unknown_excel_columns(
        list(rows[0].keys()), GLOBAL_TASK_COLUMN_ALIASES, col_index
    )
    if unknown_columns:
        warnings.append(
            f"'{sheet_name}' vərəqində əlavə sahə kimi saxlanacaq sütun(lar): "
            f"{', '.join(unknown_columns)}. "
            f"Bu sütunlar əsas sahə deyil, əlavə sahə kimi saxlanacaq."
        )

    # Collect unique assignee names for batch user creation
    unique_assignee_names = sorted({
        cell_to_str(
            get_excel_value(row, "assignee", GLOBAL_TASK_COLUMN_ALIASES, col_index)
        )
        for row in rows
        if cell_to_str(
            get_excel_value(row, "assignee", GLOBAL_TASK_COLUMN_ALIASES, col_index)
        )
    })
    user_map, users_created, users_reused, sample_created_users = (
        get_or_create_assignee_users(db, unique_assignee_names)
    )

    created_tasks = 0
    updated_tasks = 0
    no_change_count = 0
    skipped_rows = 0
    ambiguous_rows = 0
    imported_tasks = []
    custom_field_columns = set()

    for row in rows:
        title = cell_to_str(
            get_excel_value(row, "title", GLOBAL_TASK_COLUMN_ALIASES, col_index)
        )
        if not title:
            skipped_rows += 1
            warnings.append(
                f"Sətir {row['_row']}: Başlıq boşdur — bu sətir keçiriləcək."
            )
            continue

        row_label = title

        # --- Optional fields ---
        assignee_name = cell_to_str(
            get_excel_value(row, "assignee", GLOBAL_TASK_COLUMN_ALIASES, col_index)
        )
        description = cell_to_str(
            get_excel_value(row, "description", GLOBAL_TASK_COLUMN_ALIASES, col_index)
        )
        category_raw = cell_to_str(
            get_excel_value(row, "category", GLOBAL_TASK_COLUMN_ALIASES, col_index)
        )
        if category_raw:
            _auto_create_category_option(db, category_raw)
        external_id_raw = cell_to_str(
            get_excel_value(row, "external_id", GLOBAL_TASK_COLUMN_ALIASES, col_index)
        )

        # Priority mapping
        priority_raw = cell_to_str(
            get_excel_value(row, "priority", GLOBAL_TASK_COLUMN_ALIASES, col_index)
        )
        priority_norm = normalize_lookup_value(priority_raw) if priority_raw else ""
        priority = GLOBAL_PRIORITY_MAP.get(priority_norm, "medium")

        # Operational status mapping
        op_status_raw = cell_to_str(
            get_excel_value(row, "operational_status", GLOBAL_TASK_COLUMN_ALIASES, col_index)
        )
        op_status_norm = normalize_lookup_value(op_status_raw) if op_status_raw else ""
        op_status = GLOBAL_OPERATIONAL_STATUS_MAP.get(op_status_norm, "new")
        derived_status = GLOBAL_STATUS_FROM_OP_STATUS.get(op_status, "todo")

        # Deadline
        deadline_raw = get_excel_value(
            row, "deadline", GLOBAL_TASK_COLUMN_ALIASES, col_index
        )
        deadline_date = parse_excel_date(
            deadline_raw, row_label, "Son tarix", warnings, sheet_name=sheet_name
        )
        deadline_dt = (
            datetime.combine(deadline_date, time.min) if deadline_date is not None else None
        )

        # Start date
        start_date_raw = get_excel_value(
            row, "start_date", GLOBAL_TASK_COLUMN_ALIASES, col_index
        )
        start_date_date = parse_excel_date(
            start_date_raw, row_label, "Tarix", warnings, sheet_name=sheet_name
        )

        # Kanban
        kanban_raw = cell_to_str(
            get_excel_value(row, "kanban", GLOBAL_TASK_COLUMN_ALIASES, col_index)
        )
        kanban_enabled = parse_yes_no(kanban_raw)

        # Resolve assignee
        assigned_user = user_map.get(assignee_name) if assignee_name else None

        # Collect unknown columns into custom_fields
        row_custom_fields = _build_custom_fields_from_row(
            row, col_index, GLOBAL_TASK_COLUMN_ALIASES
        )
        for key in row_custom_fields:
            custom_field_columns.add(key)

        # --- Match existing task ---
        match = _match_task_for_import(
            db, external_id_raw, title, assignee_name, deadline_date
        )

        if match["action"] == "ambiguous":
            ambiguous_rows += 1
            warnings.extend(match["warnings"])
            continue

        if match["action"] == "update":
            existing = match["matched_task"]
            assigned_uid = assigned_user.id if assigned_user else None

            # Detect no-change: compare before updating
            changed = _has_row_changes(
                existing, title, description, category_raw, priority,
                op_status, derived_status, kanban_enabled, deadline_dt,
                start_date_date, external_id_raw, assigned_uid,
                row_custom_fields,
            )

            if not changed:
                no_change_count += 1
                imported_tasks.append({
                    "id": str(existing.id),
                    "title": existing.title,
                    "assignee": assignee_name,
                    "priority": existing.priority,
                    "operational_status": existing.operational_status,
                    "status": existing.status,
                    "deadline": deadline_date.isoformat() if deadline_date else None,
                    "custom_fields": existing.custom_fields or {},
                    "action": "no_change",
                    "match_reason": match["match_reason"],
                })
                continue

            # Update only fields present and non-empty in Excel
            if title:
                existing.title = title
            if description:
                existing.description = description
            if category_raw:
                existing.category = category_raw
                # Phase 8E.1: Do NOT write direction_program — it is now read-only legacy fallback.
            if priority_raw:
                existing.priority = priority
            if op_status_raw:
                existing.operational_status = op_status
                existing.status = derived_status
            if kanban_enabled is not None:
                existing.kanban_enabled = kanban_enabled
            if deadline_dt is not None:
                existing.deadline = deadline_dt
            if start_date_date is not None:
                existing.start_date = start_date_date
            if assigned_user:
                existing.assigned_to = assigned_user.id
            if external_id_raw:
                existing.external_id = external_id_raw
            # Merge custom_fields: existing keys preserved, Excel keys overwrite
            if row_custom_fields:
                merged = dict(existing.custom_fields or {})
                merged.update(row_custom_fields)
                existing.custom_fields = merged

            updated_tasks += 1
            imported_tasks.append({
                "id": str(existing.id),
                "title": existing.title,
                "assignee": assignee_name,
                "priority": existing.priority,
                "operational_status": existing.operational_status,
                "status": existing.status,
                "deadline": deadline_date.isoformat() if deadline_date else None,
                "custom_fields": row_custom_fields,
                "action": "update",
                "match_reason": match["match_reason"],
            })
            continue

        # --- Create new task ---
        task = models.Task(
            title=title,
            description=description,
            task_type="one_time",
            status=derived_status,
            operational_status=op_status,
            priority=priority,
            kanban_enabled=kanban_enabled if kanban_enabled is not None else True,
            frequency="none",
            reporting_required=False,
            category=category_raw,
            # Phase 8E.1: Do NOT write direction_program — it is now read-only legacy fallback.
            start_date=start_date_date,
            external_id=external_id_raw,
            deadline=deadline_dt,
            project_id=None,
            client_id=None,
            created_by=current_user.id,
            assigned_to=assigned_user.id if assigned_user else None,
            custom_fields=row_custom_fields,
        )
        db.add(task)
        db.flush()

        created_tasks += 1
        imported_tasks.append({
            "id": str(task.id),
            "title": title,
            "assignee": assignee_name,
            "priority": priority,
            "operational_status": op_status,
            "status": derived_status,
            "deadline": deadline_date.isoformat() if deadline_date else None,
            "custom_fields": row_custom_fields,
            "action": "create",
            "match_reason": match["match_reason"],
        })

    # Activity log
    create_activity_log(
        db=db,
        user_id=current_user.id,
        action="GLOBAL_EXCEL_IMPORT",
        entity_type="excel_import",
        entity_id=None,
        old_data=None,
        new_data={
            "created_tasks": created_tasks,
            "updated_tasks": updated_tasks,
            "no_change_count": no_change_count,
            "skipped_rows": skipped_rows,
            "ambiguous_rows": ambiguous_rows,
            "users_created": users_created,
            "users_reused": users_reused,
        },
    )

    return {
        "mode": "global",
        "created_tasks": created_tasks,
        "updated_tasks": updated_tasks,
        "no_change_count": no_change_count,
        "skipped_rows": skipped_rows,
        "ambiguous_rows": ambiguous_rows,
        "tasks": imported_tasks,
        "warnings": warnings,
        "unknown_columns": unknown_columns,
        "custom_field_columns": sorted(custom_field_columns),
    }


# ---------------------------------------------------------------------------
# Global Excel endpoints
# ---------------------------------------------------------------------------


@router.post("/global/preview")
def preview_global_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin_or_manager),
):
    """Preview a global (task-only) Excel workbook without writing to DB."""
    safe_name = Path(file.filename or "").name
    extension = Path(safe_name).suffix.lower()
    if extension not in (".xlsx", ".xls"):
        raise HTTPException(
            status_code=400,
            detail="Yalnız .xlsx və .xls faylları dəstəklənir",
        )

    validate_excel_file_size(file)
    contents = file.file.read()
    result = parse_global_preview(contents, safe_name)
    _apply_preview_matching(result, db)
    return result


@router.post("/global/import")
def import_global_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    """Import tasks from a global (task-only) Excel workbook."""
    safe_name = Path(file.filename or "").name
    extension = Path(safe_name).suffix.lower()
    if extension not in (".xlsx", ".xls"):
        raise HTTPException(
            status_code=400,
            detail="Yalnız .xlsx və .xls faylları dəstəklənir",
        )

    validate_excel_file_size(file)
    contents = file.file.read()

    try:
        result = execute_global_import(contents, db, current_user)
        db.commit()
        return result
    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        logger.exception("Global Excel import failed")
        raise_safe_excel_import_error()
        raise HTTPException(
            status_code=500,
            detail=f"Excel import xətası: {str(exc)}",
        ) from exc


@router.get("/global/template")
def download_global_template(
    current_user: models.User = Depends(require_admin_or_manager),
):
    """Download a minimal global Excel template for task import."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tapşırıqlar"

    headers = ["Xarici ID", "Tapşırıq", "Təsvir", "İcraçı", "Prioritet", "Status", "Başlama tarixi", "Son tarix", "Filial", "Müqavilə nömrəsi"]
    ws.append(headers)

    # Example row
    ws.append(["ID-001", "Nümunə tapşırıq", "Təsvir nümunəsi", "İcraçı Adı", "Orta", "Yeni", "", "", "Bakı", "CN-102"])

    # Basic styling
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Column widths
    col_widths = [14, 28, 30, 22, 12, 16, 16, 16, 18, 22]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = "task-manager-global-template.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
