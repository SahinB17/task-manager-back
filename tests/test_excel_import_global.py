"""
Automated tests for the global Excel import system.

Covers:
  1. Single-sheet English headers preview
  2. Azerbaijani headers preview
  3. Import creates task with project_id = null
  4. Missing title column returns 400 with required_missing
  5. Row with empty title skipped with warning, other rows import fine
  6. Unknown column appears in warnings but does not block
  7. Removed /excel/sxa workflow returns 404
"""

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_global_workbook_bytes(headers, rows, sheet_name="Sheet1"):
    """Build an in-memory .xlsx workbook with a single task sheet.

    Returns BytesIO positioned at start ready for upload.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _upload_global_preview(client, auth_headers_dict, workbook_bytes, filename="test.xlsx"):
    """POST workbook to /excel/global/preview, return response."""
    return client.post(
        "/excel/global/preview",
        files={
            "file": (
                filename,
                workbook_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=auth_headers_dict,
    )


def _upload_global_import(client, auth_headers_dict, workbook_bytes, filename="test.xlsx"):
    """POST workbook to /excel/global/import, return response."""
    return client.post(
        "/excel/global/import",
        files={
            "file": (
                filename,
                workbook_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=auth_headers_dict,
    )


def test_start_date_migration_sql_exists_and_backfills_legacy_column():
    migration_path = Path(__file__).resolve().parents[1] / "sql" / "add_task_start_date.sql"

    assert migration_path.exists()
    sql = migration_path.read_text(encoding="utf-8").lower()
    assert "add column if not exists start_date date" in sql
    assert "set start_date = sxa_start_date" in sql
    assert "where start_date is null" in sql
    assert "sxa_start_date is not null" in sql
    assert "drop column" not in sql


def test_category_migration_sql_exists_and_backfills_legacy_column():
    migration_path = Path(__file__).resolve().parents[1] / "sql" / "add_task_category.sql"

    assert migration_path.exists()
    sql = migration_path.read_text(encoding="utf-8").lower()
    assert "add column if not exists category" in sql
    assert "set category = direction_program" in sql
    assert "where category is null" in sql
    assert "direction_program is not null" in sql
    assert "drop column" not in sql


def test_legacy_sxa_preview_route_removed(client, auth_headers, factories):
    factories["create_user"](
        full_name="Admin Removed SXA Preview",
        email="admin-removed-sxa-preview@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-removed-sxa-preview@test.local", "adminpass")
    wb_bytes = _make_global_workbook_bytes(["Title"], [["Removed route check"]])

    response = client.post(
        "/excel/sxa/preview",
        files={
            "file": (
                "legacy-sxa.xlsx",
                wb_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )

    assert response.status_code == 404


def test_legacy_sxa_import_route_removed(client, auth_headers, factories):
    factories["create_user"](
        full_name="Admin Removed SXA Import",
        email="admin-removed-sxa-import@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-removed-sxa-import@test.local", "adminpass")
    wb_bytes = _make_global_workbook_bytes(["Title"], [["Removed route check"]])

    response = client.post(
        "/excel/sxa/import",
        files={
            "file": (
                "legacy-sxa.xlsx",
                wb_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )

    assert response.status_code == 404


def test_global_export_workbook_can_be_reimported(
    client,
    auth_headers,
    factories,
    db_session,
    app_modules,
):
    """Exported task workbook uses global headers and can be sent to global import."""
    from datetime import date

    models = app_modules["models"]
    admin = factories["create_user"](
        full_name="Admin Export",
        email="admin-export-global@test.local",
        role_id=1,
        password="adminpass",
    )
    worker = factories["create_user"](
        full_name="Export Worker",
        email="export-worker@test.local",
        role_id=3,
        password="workerpass",
    )
    task = models.Task(
        title="Global export reimport",
        description="Reusable export row",
        created_by=admin.id,
        assigned_to=worker.id,
        external_id="EXP-REIMPORT-1",
        sxa_task_id="SXA-FALLBACK-1",
        priority="high",
        operational_status="active",
        status="in_progress",
        start_date=date(2026, 7, 1),
        current_deadline=date(2026, 7, 15),
        category="Audit",
        direction_program="Audit",
        kanban_enabled=True,
        custom_fields={"Filial": "Bakı", "Müqavilə nömrəsi": "CN-102"},
    )
    db_session.add(task)
    db_session.commit()

    headers = auth_headers("admin-export-global@test.local", "adminpass")
    export_response = client.get("/excel/tasks/export", headers=headers)
    assert export_response.status_code == 200

    workbook = load_workbook(BytesIO(export_response.content), data_only=True)
    try:
        rows = list(workbook["Tapşırıqlar"].iter_rows(values_only=True))
        assert rows[0][:12] == (
            "Xarici ID",
            "Tapşırıq",
            "Təsvir",
            "İcraçı",
            "Prioritet",
            "Status",
            "Başlama tarixi",
            "Son tarix",
            "Kateqoriya",
            "Kanban",
            "Filial",
            "Müqavilə nömrəsi",
        )
        assert rows[1][0] == "EXP-REIMPORT-1"
        assert rows[1][6].date() == date(2026, 7, 1)
        assert rows[1][10:12] == ("Bakı", "CN-102")
    finally:
        workbook.close()

    import_response = _upload_global_import(
        client,
        headers,
        BytesIO(export_response.content),
        filename="global-export.xlsx",
    )
    assert import_response.status_code == 200, import_response.text
    assert db_session.query(models.Task).filter(
        models.Task.external_id == "EXP-REIMPORT-1"
    ).count() == 1
    db_session.refresh(task)
    assert task.start_date == date(2026, 7, 1)
    assert task.custom_fields == {"Filial": "Bakı", "Müqavilə nömrəsi": "CN-102"}


def test_global_template_download_works(client, auth_headers, factories):
    factories["create_user"](
        full_name="Admin Template",
        email="admin-template-global@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-template-global@test.local", "adminpass")

    response = client.get("/excel/global/template", headers=headers)

    assert response.status_code == 200
    assert "task-manager-global-template.xlsx" in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content), data_only=True)
    try:
        assert workbook.sheetnames == ["Tapşırıqlar"]
        rows = list(workbook["Tapşırıqlar"].iter_rows(values_only=True))
        assert rows[0] == (
            "Xarici ID",
            "Tapşırıq",
            "Təsvir",
            "İcraçı",
            "Prioritet",
            "Status",
            "Başlama tarixi",
            "Son tarix",
            "Filial",
            "Müqavilə nömrəsi",
        )
        assert rows[1][0] == "ID-001"
        assert rows[1][1] == "Nümunə tapşırıq"
    finally:
        workbook.close()


# ---------------------------------------------------------------------------
# 1. Single-sheet Sheet1 with English headers previews successfully
# ---------------------------------------------------------------------------


def test_global_preview_single_sheet_english_headers(client, auth_headers, factories):
    """Sheet1 with Title, Assigned to, Priority, Deadline, Status returns 200."""
    factories["create_user"](
        full_name="Admin Global",
        email="admin-global-eng@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-global-eng@test.local", "adminpass")

    task_headers = ["Title", "Assigned to", "Priority", "Deadline", "Status"]
    task_rows = [
        ["Simple English Task", "Worker One", "High", None, "Active"],
    ]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_preview(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Expected 200, got {resp.status_code}: {resp.text}"
    data = resp.json()
    assert data["mode"] == "global"
    assert data["tasks_count"] == 1
    assert len(data["tasks"]) == 1
    assert data["tasks"][0]["title"] == "Simple English Task"
    assert data["tasks"][0]["assignee"] == "Worker One"
    assert data["tasks"][0]["priority"] == "High"
    assert data["tasks"][0]["operational_status"] == "Active"
    # required_missing should be empty since title is present
    assert data["required_missing"] == []


# ---------------------------------------------------------------------------
# 2. Azerbaijani headers preview successfully
# ---------------------------------------------------------------------------


def test_global_preview_azerbaijani_headers(client, auth_headers, factories):
    """Azerbaijani column headers Tapşırıq, İcraçı, Prioritet, Son tarix, Status preview 200."""
    factories["create_user"](
        full_name="Admin Global",
        email="admin-global-az@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-global-az@test.local", "adminpass")

    task_headers = ["Tapşırıq", "İcraçı", "Prioritet", "Son tarix", "Status"]
    task_rows = [
        ["Azərbaycan Tapşırığı", "İşçi Bir", "Yüksək", None, "İcradadır"],
    ]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_preview(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Expected 200, got {resp.status_code}: {resp.text}"
    data = resp.json()
    assert data["mode"] == "global"
    assert data["tasks_count"] == 1
    assert data["tasks"][0]["title"] == "Azərbaycan Tapşırığı"
    assert data["tasks"][0]["assignee"] == "İşçi Bir"
    assert data["tasks"][0]["priority"] == "Yüksək"
    assert data["tasks"][0]["operational_status"] == "İcradadır"


# ---------------------------------------------------------------------------
# 3. Import creates task with project_id = null
# ---------------------------------------------------------------------------


def test_global_import_creates_task_project_id_null(
    client, auth_headers, factories, db_session, app_modules
):
    """Imported task has project_id = null (no project association)."""
    models = app_modules["models"]
    admin = factories["create_user"](
        full_name="Admin Global",
        email="admin-global-import@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-global-import@test.local", "adminpass")

    task_headers = ["Tapşırıq", "İcraçı"]
    task_rows = [
        ["Project Null Task", "Worker Name"],
    ]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_import(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Import failed: {resp.text}"
    data = resp.json()
    assert data["mode"] == "global"
    assert data["created_tasks"] == 1
    assert data["updated_tasks"] == 0
    assert data["skipped_rows"] == 0

    # Verify task exists in DB with project_id = null
    task = (
        db_session.query(models.Task)
        .filter(models.Task.title == "Project Null Task")
        .first()
    )
    assert task is not None
    assert task.project_id is None
    assert task.client_id is None
    assert task.task_type == "one_time"
    assert task.frequency == "none"
    assert task.reporting_required == False  # noqa: E712
    assert task.kanban_enabled == True  # noqa: E712


def test_global_import_writes_start_date_to_neutral_column(
    client, auth_headers, factories, db_session, app_modules
):
    """Imported start date writes to tasks.start_date, not legacy sxa_start_date."""
    from datetime import date

    models = app_modules["models"]
    factories["create_user"](
        full_name="Admin Start Date",
        email="admin-start-date-import@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-start-date-import@test.local", "adminpass")

    task_headers = ["Xarici ID", "Tapşırıq", "Başlama tarixi"]
    task_rows = [["START-001", "Start Date Import Task", "2026-02-03"]]

    resp = _upload_global_import(
        client,
        headers,
        _make_global_workbook_bytes(task_headers, task_rows),
    )

    assert resp.status_code == 200, f"Import failed: {resp.text}"
    task = (
        db_session.query(models.Task)
        .filter(models.Task.external_id == "START-001")
        .first()
    )
    assert task is not None
    assert task.start_date == date(2026, 2, 3)
    assert task.sxa_start_date is None


# ---------------------------------------------------------------------------
# 4. Missing title column returns clear 400 with required_missing
# ---------------------------------------------------------------------------


def test_global_preview_missing_title_column_400(client, auth_headers, factories):
    """Workbook with no title column returns 400 or required_missing: ['title']."""
    factories["create_user"](
        full_name="Admin Global",
        email="admin-global-missing@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-global-missing@test.local", "adminpass")

    # No recognized title column
    task_headers = ["Priority", "Deadline", "Notes"]
    task_rows = [
        ["High", "2026-07-15", "Some note"],
    ]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_preview(client, headers, wb_bytes)

    # Accept either 400 (sheet not found) or 200 with required_missing
    if resp.status_code == 400:
        detail = resp.json()["detail"]
        if isinstance(detail, dict):
            assert "title" in detail.get("required_missing", []) or "Başlıq" in detail.get("message", "")
        else:
            assert "title" in str(detail).lower() or "tapşırıq" in str(detail).lower()
    elif resp.status_code == 200:
        data = resp.json()
        assert "title" in data["required_missing"], f"Expected title in required_missing: {data}"
        assert data["tasks_count"] == 0
    else:
        raise AssertionError(f"Unexpected status {resp.status_code}: {resp.text}")


# ---------------------------------------------------------------------------
# 5. Row with empty title skipped, other rows import fine
# ---------------------------------------------------------------------------


def test_global_import_empty_title_skipped_others_imported(
    client, auth_headers, factories, db_session, app_modules
):
    """Row with empty title is skipped with warning; valid rows are imported."""
    models = app_modules["models"]
    admin = factories["create_user"](
        full_name="Admin Global",
        email="admin-global-skip@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-global-skip@test.local", "adminpass")

    task_headers = ["Tapşırıq", "İcraçı", "Prioritet"]
    task_rows = [
        ["Valid Task One", "Worker A", "High"],
        ["", "Worker B", "Medium"],           # empty title → should be skipped
        ["Valid Task Two", "Worker C", "Low"],
    ]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_import(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Import failed: {resp.text}"
    data = resp.json()
    assert data["created_tasks"] == 2
    assert data["skipped_rows"] == 1

    # Verify warnings mention the empty title row
    skip_warning = next(
        (w for w in data["warnings"] if "başlıq" in w.lower() or "Başlıq" in w),
        None,
    )
    assert skip_warning is not None, f"No skip warning found in: {data['warnings']}"

    # Verify only valid tasks exist in DB
    t1 = (
        db_session.query(models.Task)
        .filter(models.Task.title == "Valid Task One")
        .first()
    )
    t2 = (
        db_session.query(models.Task)
        .filter(models.Task.title == "Valid Task Two")
        .first()
    )
    assert t1 is not None
    assert t2 is not None

    # Verify no empty-title task was created
    empty_count = (
        db_session.query(models.Task)
        .filter(models.Task.title == "")
        .count()
    )
    assert empty_count == 0


# ---------------------------------------------------------------------------
# 6. Unknown column appears in warnings but does not block import
# ---------------------------------------------------------------------------


def test_global_preview_unknown_column_warning_does_not_block(client, auth_headers, factories):
    """Unknown column name appears in unknown_columns and warnings, returns 200."""
    factories["create_user"](
        full_name="Admin Global",
        email="admin-global-unknown@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-global-unknown@test.local", "adminpass")

    task_headers = ["Tapşırıq", "UnknownColumnXYZ", "Prioritet"]
    task_rows = [
        ["Test Task", "Some Value", "High"],
    ]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_preview(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Expected 200, got {resp.status_code}: {resp.text}"
    data = resp.json()
    assert data["tasks_count"] == 1  # task still recognized
    assert "UnknownColumnXYZ" in data["unknown_columns"]

    # A warning about unknown columns should exist
    unknown_warning = next(
        (w for w in data["warnings"] if "UnknownColumnXYZ" in w or "tanınmayan" in w.lower()),
        None,
    )
    assert unknown_warning is not None, f"No unknown column warning in: {data['warnings']}"


# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------
# Phase 6B: Unknown columns → custom_fields
# ---------------------------------------------------------------------------


def test_global_preview_unknown_columns_become_custom_fields(client, auth_headers, factories):
    """Preview: unknown columns appear in custom_fields per task."""
    factories["create_user"](
        full_name="Admin CF",
        email="admin-cf-excel1@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-cf-excel1@test.local", "adminpass")

    task_headers = ["Title", "Branch", "Contract No"]
    task_rows = [
        ["Test Task CF", "Bakı", "CN-102"],
    ]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_preview(client, headers, wb_bytes)

    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["tasks_count"] == 1
    task = data["tasks"][0]
    assert task["title"] == "Test Task CF"
    assert task["custom_fields"] == {"Branch": "Bakı", "Contract No": "CN-102"}
    assert "Branch" in data["custom_field_columns"]
    assert "Contract No" in data["custom_field_columns"]


def test_global_import_persists_custom_fields_in_db(
    client, auth_headers, factories, db_session, app_modules,
):
    """Import: unknown columns persist as custom_fields in DB."""
    models = app_modules["models"]
    factories["create_user"](
        full_name="Admin CF",
        email="admin-cf-excel2@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-cf-excel2@test.local", "adminpass")

    task_headers = ["Tapşırıq", "Filial", "Müqavilə nömrəsi"]
    task_rows = [
        ["Import CF Task", "Bakı", "AZ-55"],
    ]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_import(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Import failed: {resp.text}"
    data = resp.json()
    assert data["created_tasks"] == 1
    assert data["custom_field_columns"] == ["Filial", "Müqavilə nömrəsi"]
    # Verify in DB
    task = (
        db_session.query(models.Task)
        .filter(models.Task.title == "Import CF Task")
        .first()
    )
    assert task is not None
    assert task.custom_fields == {"Filial": "Bakı", "Müqavilə nömrəsi": "AZ-55"}


def test_global_preview_empty_custom_field_values_skipped(client, auth_headers, factories):
    """Empty / null unknown column values are not added to custom_fields."""
    factories["create_user"](
        full_name="Admin CF",
        email="admin-cf-excel3@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-cf-excel3@test.local", "adminpass")

    task_headers = ["Title", "Priority", "Branch", "Contract No"]
    task_rows = [
        ["Only Branch", "High", "Bakı", None],       # null → skipped
        ["Only Contract", "Medium", None, "CN-99"],    # null → skipped
    ]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_preview(client, headers, wb_bytes)

    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["tasks_count"] == 2

    # Task 1: only Branch
    t1 = data["tasks"][0]
    assert t1["title"] == "Only Branch"
    assert t1["custom_fields"] == {"Branch": "Bakı"}
    assert "Contract No" not in t1["custom_fields"]

    # Task 2: only Contract No
    t2 = data["tasks"][1]
    assert t2["title"] == "Only Contract"
    assert t2["custom_fields"] == {"Contract No": "CN-99"}
    assert "Branch" not in t2["custom_fields"]


def test_global_preview_known_columns_not_in_custom_fields(client, auth_headers, factories):
    """Known columns (Title, Priority) must NOT appear in custom_fields."""
    factories["create_user"](
        full_name="Admin CF",
        email="admin-cf-excel4@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-cf-excel4@test.local", "adminpass")

    task_headers = ["Tapşırıq", "Prioritet", "İcraçı", "Son tarix", "Status", "Filial"]
    task_rows = [
        ["Known Test", "Yüksək", "Worker A", None, "İcradadır", "Gəncə"],
    ]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_preview(client, headers, wb_bytes)

    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["tasks_count"] == 1
    task = data["tasks"][0]

    # Known fields should be resolved as core fields, not in custom_fields
    assert task["title"] == "Known Test"
    assert task["priority"] == "Yüksək"
    assert task["assignee"] == "Worker A"
    assert task["operational_status"] == "İcradadır"

    # Only Filial should be in custom_fields
    assert task["custom_fields"] == {"Filial": "Gəncə"}
    # Known column names must NOT be in custom_fields keys
    for known in ["Tapşırıq", "Prioritet", "İcraçı", "Son tarix", "Status"]:
        assert known not in task["custom_fields"], f"'{known}' should not be in custom_fields"


def test_global_preview_unknown_columns_warning_language_soft(client, auth_headers, factories):
    """Unknown column warning uses softer language (əlavə sahə, not tanınmayan)."""
    factories["create_user"](
        full_name="Admin CF",
        email="admin-cf-excel5@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-cf-excel5@test.local", "adminpass")

    task_headers = ["Title", "UnknownCol"]
    task_rows = [["Soft Warning Test", "value"]]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_preview(client, headers, wb_bytes)

    assert resp.status_code == 200, resp.text
    data = resp.json()

    # Warning should say "əlavə sahə kimi saxlanacaq" not "tanınmayan sütunlar" (the old scary wording)
    all_warnings = " ".join(data["warnings"])
    assert "əlavə sahə" in all_warnings.lower()


def test_global_import_azerbaijani_custom_fields(
    client, auth_headers, factories, db_session, app_modules,
):
    """Azerbaijani unknown columns work as custom_fields."""
    models = app_modules["models"]
    factories["create_user"](
        full_name="Admin CF",
        email="admin-cf-excel6@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-cf-excel6@test.local", "adminpass")

    task_headers = ["Tapşırıq", "Filial", "Müqavilə nömrəsi", "Ödəniş statusu"]
    task_rows = [
        ["Az CF Task", "Sumqayıt", "AZ-55", "Ödənilib"],
    ]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_import(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Import failed: {resp.text}"
    data = resp.json()
    assert data["created_tasks"] == 1

    task = (
        db_session.query(models.Task)
        .filter(models.Task.title == "Az CF Task")
        .first()
    )
    assert task is not None
    assert task.custom_fields == {
        "Filial": "Sumqayıt",
        "Müqavilə nömrəsi": "AZ-55",
        "Ödəniş statusu": "Ödənilib",
    }


# ---------------------------------------------------------------------------
# Phase 6C: Upsert / dedup via external_id
# ---------------------------------------------------------------------------


def test_global_import_new_external_id_creates_task(
    client, auth_headers, factories, db_session, app_modules,
):
    """Import with new external_id creates a task (no duplicate)."""
    models = app_modules["models"]
    factories["create_user"](
        full_name="Admin Upsert",
        email="admin-upsert1@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-upsert1@test.local", "adminpass")

    task_headers = ["ID", "Tapşırıq", "Prioritet"]
    task_rows = [["EXT-001", "New External ID Task", "High"]]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_import(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Import failed: {resp.text}"
    data = resp.json()
    assert data["created_tasks"] == 1
    assert data["updated_tasks"] == 0

    task = (
        db_session.query(models.Task)
        .filter(models.Task.external_id == "EXT-001")
        .first()
    )
    assert task is not None
    assert task.title == "New External ID Task"
    assert task.external_id == "EXT-001"


def test_global_import_existing_external_id_updates_task(
    client, auth_headers, factories, db_session, app_modules,
):
    """Import with existing external_id updates task, does not create duplicate."""
    models = app_modules["models"]
    admin = factories["create_user"](
        full_name="Admin Upsert",
        email="admin-upsert2@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-upsert2@test.local", "adminpass")

    # Create a task first via API
    create_resp = client.post(
        "/tasks/",
        json={"title": "Original Title", "external_id": "EXT-002", "priority": "low"},
        headers=headers,
    )
    assert create_resp.status_code == 200

    # Import with same external_id but updated fields
    task_headers = ["ID", "Tapşırıq", "Prioritet"]
    task_rows = [["EXT-002", "Updated Title", "High"]]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_import(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Import failed: {resp.text}"
    data = resp.json()
    assert data["updated_tasks"] == 1
    assert data["created_tasks"] == 0

    # Verify task was updated, not duplicated
    tasks = (
        db_session.query(models.Task)
        .filter(models.Task.external_id == "EXT-002")
        .all()
    )
    assert len(tasks) == 1, f"Expected 1 task, got {len(tasks)}"
    assert tasks[0].title == "Updated Title"
    assert tasks[0].priority == "high"


def test_global_import_same_file_twice_no_duplicate(
    client, auth_headers, factories, db_session, app_modules,
):
    """Import same file twice with external_id does not create duplicates."""
    models = app_modules["models"]
    factories["create_user"](
        full_name="Admin Upsert",
        email="admin-upsert3@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-upsert3@test.local", "adminpass")

    task_headers = ["ID", "Tapşırıq", "Prioritet"]
    task_rows = [["EXT-003", "Dedup Task", "Medium"]]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)

    # First import
    r1 = _upload_global_import(client, headers, wb_bytes)
    assert r1.status_code == 200
    assert r1.json()["created_tasks"] == 1

    # Second import — same file
    wb_bytes2 = _make_global_workbook_bytes(task_headers, task_rows)
    r2 = _upload_global_import(client, headers, wb_bytes2)
    assert r2.status_code == 200
    # Second import: identical row → no_change, not update
    assert r2.json()["no_change_count"] == 1
    assert r2.json()["created_tasks"] == 0
    assert r2.json()["updated_tasks"] == 0

    # Verify only one task exists
    tasks = (
        db_session.query(models.Task)
        .filter(models.Task.external_id == "EXT-003")
        .all()
    )
    assert len(tasks) == 1


def test_global_import_update_preserves_existing_fields(
    client, auth_headers, factories, db_session, app_modules,
):
    """Update does not erase existing fields when Excel cells are empty."""
    models = app_modules["models"]
    admin = factories["create_user"](
        full_name="Admin Upsert",
        email="admin-upsert4@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-upsert4@test.local", "adminpass")

    # Create task with description
    create_resp = client.post(
        "/tasks/",
        json={
            "title": "Preserve Desc",
            "description": "original description",
            "external_id": "EXT-004",
            "priority": "low",
            "direction_program": "IT",
            "start_date": "2026-02-01",
        },
        headers=headers,
    )
    assert create_resp.status_code == 200

    # Import with external_id — only update priority, leave description/category empty
    task_headers = ["ID", "Tapşırıq", "Prioritet"]
    task_rows = [["EXT-004", "Preserve Desc Updated", "High"]]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_import(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Import failed: {resp.text}"
    data = resp.json()
    assert data["updated_tasks"] == 1

    task = (
        db_session.query(models.Task)
        .filter(models.Task.external_id == "EXT-004")
        .first()
    )
    assert task.title == "Preserve Desc Updated"
    assert task.priority == "high"
    # Description and category should be preserved (not overwritten with empty)
    assert task.description == "original description"
    assert task.category == "IT"
    assert task.direction_program is None  # Phase 8E.1: direction_program is now read-only legacy fallback.
    assert task.start_date.isoformat() == "2026-02-01"


def test_global_import_category_auto_creates_option(
    client, auth_headers, factories, db_session, app_modules,
):
    """Category column import creates a matching category option."""
    models = app_modules["models"]
    factories["create_user"](
        full_name="Admin Category",
        email="admin-category-import@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-category-import@test.local", "adminpass")

    task_headers = ["Xarici ID", "Tapşırıq", "Kateqoriya"]
    task_rows = [["CAT-001", "Category Import Task", "Yeni Kateqoriya"]]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_import(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Import failed: {resp.text}"
    option = (
        db_session.query(models.CategoryOption)
        .filter(models.CategoryOption.name == "Yeni Kateqoriya")
        .first()
    )
    assert option is not None

    task = (
        db_session.query(models.Task)
        .filter(models.Task.external_id == "CAT-001")
        .first()
    )
    assert task is not None
    assert task.category == "Yeni Kateqoriya"
    assert task.direction_program is None  # Phase 8E.1: direction_program is now read-only legacy fallback.


def test_global_import_custom_fields_merge_on_update(
    client, auth_headers, factories, db_session, app_modules,
):
    """custom_fields merge on update: existing keys preserved, Excel keys overwrite."""
    models = app_modules["models"]
    factories["create_user"](
        full_name="Admin Upsert",
        email="admin-upsert5@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-upsert5@test.local", "adminpass")

    # First import with Branch: Bakı
    task_headers1 = ["ID", "Tapşırıq", "Branch"]
    task_rows1 = [["EXT-005", "Merge CF Task", "Bakı"]]
    r1 = _upload_global_import(client, headers, _make_global_workbook_bytes(task_headers1, task_rows1))
    assert r1.status_code == 200

    # Manually add an "Old" key to custom_fields (simulating external modification)
    task = (
        db_session.query(models.Task)
        .filter(models.Task.external_id == "EXT-005")
        .first()
    )
    task.custom_fields = {"Old": "keep", "Branch": "Bakı"}
    db_session.commit()

    # Second import: Branch → Gəncə, new Contract No column
    task_headers2 = ["ID", "Tapşırıq", "Branch", "Contract No"]
    task_rows2 = [["EXT-005", "Merge CF Task", "Gəncə", "CN-1"]]
    r2 = _upload_global_import(client, headers, _make_global_workbook_bytes(task_headers2, task_rows2))
    assert r2.status_code == 200

    # Reload
    db_session.expire_all()
    task = (
        db_session.query(models.Task)
        .filter(models.Task.external_id == "EXT-005")
        .first()
    )
    assert task.custom_fields == {
        "Old": "keep",
        "Branch": "Gəncə",
        "Contract No": "CN-1",
    }


def test_global_preview_returns_action_counts(
    client, auth_headers, factories,
):
    """Preview returns create_count, update_count, etc."""
    factories["create_user"](
        full_name="Admin Upsert",
        email="admin-upsert6@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-upsert6@test.local", "adminpass")

    task_headers = ["ID", "Tapşırıq", "Prioritet"]
    task_rows = [["EXT-NEW", "Preview Action Task", "Orta"]]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_preview(client, headers, wb_bytes)

    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["create_count"] == 1
    assert "update_count" in data
    assert "ambiguous_count" in data
    assert "skipped_count" in data
    # Each task should have action
    assert data["tasks"][0]["action"] == "create"
    assert data["tasks"][0]["match_reason"] == "external_id"


# ---------------------------------------------------------------------------
# Phase 6C polish: no-change detection
# ---------------------------------------------------------------------------


def test_global_preview_identical_row_returns_no_change(
    client, auth_headers, factories,
):
    """Preview: existing task with same external_id and identical row → no_change."""
    factories["create_user"](
        full_name="Admin NC",
        email="admin-nc1@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-nc1@test.local", "adminpass")

    # Create task with known external_id
    client.post(
        "/tasks/",
        json={"title": "No Change Task", "external_id": "NC-001", "priority": "medium"},
        headers=headers,
    )

    # Preview identical row
    task_headers = ["ID", "Tapşırıq", "Prioritet"]
    task_rows = [["NC-001", "No Change Task", "Orta"]]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_preview(client, headers, wb_bytes)

    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["tasks"][0]["action"] == "no_change"
    assert data["no_change_count"] == 1


def test_global_import_identical_row_does_not_update(
    client, auth_headers, factories, db_session, app_modules,
):
    """Import identical row does not change updated_at or call update."""
    models = app_modules["models"]
    factories["create_user"](
        full_name="Admin NC",
        email="admin-nc2@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-nc2@test.local", "adminpass")

    # Create task
    create_resp = client.post(
        "/tasks/",
        json={"title": "Preserve Updated At", "external_id": "NC-002", "priority": "high"},
        headers=headers,
    )
    assert create_resp.status_code == 200
    task_id = create_resp.json()["id"]

    task_before = (
        db_session.query(models.Task)
        .filter(models.Task.id == task_id)
        .first()
    )
    original_updated_at = task_before.updated_at

    # Import identical row
    task_headers = ["ID", "Tapşırıq", "Prioritet"]
    task_rows = [["NC-002", "Preserve Updated At", "Yüksək"]]

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_import(client, headers, wb_bytes)

    assert resp.status_code == 200, f"Import failed: {resp.text}"
    data = resp.json()
    assert data["updated_tasks"] == 0
    assert data["no_change_count"] == 1

    # Verify updated_at didn't change
    db_session.expire_all()
    task_after = (
        db_session.query(models.Task)
        .filter(models.Task.id == task_id)
        .first()
    )
    assert task_after.updated_at == original_updated_at


def test_global_preview_changed_title_returns_update(
    client, auth_headers, factories,
):
    """Existing task with same external_id but changed title → update."""
    factories["create_user"](
        full_name="Admin NC",
        email="admin-nc3@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-nc3@test.local", "adminpass")

    client.post(
        "/tasks/",
        json={"title": "Old Title", "external_id": "NC-003", "priority": "low"},
        headers=headers,
    )

    task_headers = ["ID", "Tapşırıq", "Prioritet"]
    task_rows = [["NC-003", "New Title", "Aşağı"]]  # same priority as "low"

    wb_bytes = _make_global_workbook_bytes(task_headers, task_rows)
    resp = _upload_global_preview(client, headers, wb_bytes)

    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["tasks"][0]["action"] == "update"
    assert data["update_count"] == 1


def test_global_preview_custom_fields_identical_returns_no_change(
    client, auth_headers, factories, db_session, app_modules,
):
    """custom_fields identical → no_change."""
    models = app_modules["models"]
    factories["create_user"](
        full_name="Admin NC",
        email="admin-nc4@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-nc4@test.local", "adminpass")

    # Create task with custom_fields via import
    task_headers1 = ["ID", "Tapşırıq", "Branch"]
    task_rows1 = [["NC-004", "CF No Change", "Bakı"]]
    r1 = _upload_global_import(
        client, headers,
        _make_global_workbook_bytes(task_headers1, task_rows1),
    )
    assert r1.status_code == 200

    # Preview identical row
    task_headers2 = ["ID", "Tapşırıq", "Branch"]
    task_rows2 = [["NC-004", "CF No Change", "Bakı"]]
    r2 = _upload_global_preview(
        client, headers,
        _make_global_workbook_bytes(task_headers2, task_rows2),
    )
    assert r2.status_code == 200, r2.text
    data = r2.json()
    assert data["tasks"][0]["action"] == "no_change"


def test_global_preview_custom_fields_changed_returns_update(
    client, auth_headers, factories,
):
    """custom_fields changed/added → update."""
    factories["create_user"](
        full_name="Admin NC",
        email="admin-nc5@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-nc5@test.local", "adminpass")

    # Create task with Branch: Bakı
    task_headers1 = ["ID", "Tapşırıq", "Branch"]
    task_rows1 = [["NC-005", "CF Changed", "Bakı"]]
    r1 = _upload_global_import(
        client, headers,
        _make_global_workbook_bytes(task_headers1, task_rows1),
    )
    assert r1.status_code == 200

    # Preview with changed Branch + new Contract No
    task_headers2 = ["ID", "Tapşırıq", "Branch", "Contract No"]
    task_rows2 = [["NC-005", "CF Changed", "Gəncə", "CN-99"]]
    r2 = _upload_global_preview(
        client, headers,
        _make_global_workbook_bytes(task_headers2, task_rows2),
    )
    assert r2.status_code == 200, r2.text
    data = r2.json()
    assert data["tasks"][0]["action"] == "update"


# ---------------------------------------------------------------------------
# Production fix: file size validation
# ---------------------------------------------------------------------------


def test_global_preview_rejects_oversized_file(client, auth_headers, factories):
    """POST /excel/global/preview rejects file over 10 MB."""
    factories["create_user"](
        full_name="Admin Size",
        email="admin-size@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-size@test.local", "adminpass")

    oversized = BytesIO(b"x" * (11 * 1024 * 1024))

    resp = client.post(
        "/excel/global/preview",
        files={
            "file": (
                "oversized.xlsx",
                oversized,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )
    assert resp.status_code == 400, f"Expected 400, got {resp.status_code}: {resp.text}"
    assert "ölçü" in resp.json()["detail"].lower() or "böyük" in resp.json()["detail"].lower()


def test_global_import_rejects_oversized_file(client, auth_headers, factories):
    """POST /excel/global/import rejects file over 10 MB."""
    factories["create_user"](
        full_name="Admin Size",
        email="admin-size2@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin-size2@test.local", "adminpass")

    oversized = BytesIO(b"x" * (11 * 1024 * 1024))

    resp = client.post(
        "/excel/global/import",
        files={
            "file": (
                "oversized.xlsx",
                oversized,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )
    assert resp.status_code == 400, f"Expected 400, got {resp.status_code}: {resp.text}"


def test_global_export_category_falls_back_to_direction_program(
    client, auth_headers, factories, db_session, app_modules,
):
    """Export uses task.category first, falls back to task.direction_program."""
    from datetime import date

    models = app_modules["models"]
    admin = factories["create_user"](
        full_name="Admin Cat Export",
        email="admin-cat-export@test.local",
        role_id=1,
        password="adminpass",
    )
    assignee = factories["create_user"](
        full_name="Worker Cat Export",
        email="worker-cat-export@test.local",
        role_id=3,
    )

    # Task with category set (primary path)
    task_with_category = models.Task(
        title="Category Set Task",
        category="Maliyyə",
        direction_program="OldDirection",
        created_by=admin.id,
        assigned_to=assignee.id,
        deadline=date(2026, 12, 31),
        status="todo",
        operational_status="new",
    )
    # Task with only direction_program (fallback path)
    task_with_legacy = models.Task(
        title="Legacy Only Task",
        direction_program="LegacyDirection",
        created_by=admin.id,
        assigned_to=assignee.id,
        deadline=date(2026, 12, 31),
        status="todo",
        operational_status="new",
    )
    db_session.add_all([task_with_category, task_with_legacy])
    db_session.commit()

    headers = auth_headers("admin-cat-export@test.local", "adminpass")
    response = client.get("/excel/tasks/export", headers=headers)
    assert response.status_code == 200

    workbook = load_workbook(BytesIO(response.content), data_only=True)
    try:
        rows = list(workbook["Tapşırıqlar"].iter_rows(values_only=True))
        # Find rows by title
        row_category = next(r for r in rows if r[1] == "Category Set Task")
        row_legacy = next(r for r in rows if r[1] == "Legacy Only Task")

        # Row with category set should export category value
        assert row_category[8] == "Maliyyə"
        # Row with only direction_program should fall back to it
        assert row_legacy[8] == "LegacyDirection"
    finally:
        workbook.close()
