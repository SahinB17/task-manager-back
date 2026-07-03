import json
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook


def _make_function_response(*, response_id, tool_name, arguments, call_id="call_1"):
    return SimpleNamespace(
        id=response_id,
        output=[
            SimpleNamespace(
                type="function_call",
                name=tool_name,
                call_id=call_id,
                arguments=json.dumps(arguments),
            )
        ],
        output_text="",
    )


def _make_text_response(*, response_id, text):
    return SimpleNamespace(
        id=response_id,
        output=[],
        output_text=text,
    )


class _FakeResponsesAPI:
    def __init__(self, payloads):
        self._payloads = list(payloads)

    def create(self, **kwargs):
        if not self._payloads:
            raise AssertionError(f"Unexpected OpenAI call: {kwargs}")
        return self._payloads.pop(0)


class _FakeOpenAIClient:
    def __init__(self, payloads):
        self.responses = _FakeResponsesAPI(payloads)


def _make_global_workbook_bytes(headers, rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def test_auth_login_and_me(client, auth_headers, factories):
    factories["create_user"](
        full_name="Admin User",
        email="admin@test.local",
        role_id=1,
        password="adminpass",
    )

    login_response = client.post(
        "/auth/login",
        data={"username": "admin@test.local", "password": "adminpass"},
    )
    assert login_response.status_code == 200
    assert login_response.json()["token_type"] == "bearer"
    assert login_response.json()["user"]["must_change_password"] is False

    me_response = client.get("/auth/me", headers=auth_headers("admin@test.local", "adminpass"))
    assert me_response.status_code == 200
    assert me_response.json()["email"] == "admin@test.local"
    assert me_response.json()["role_id"] == 1


def test_manager_can_create_employee_but_not_admin(client, auth_headers, factories):
    factories["create_user"](
        full_name="Manager User",
        email="manager@test.local",
        role_id=2,
        password="managerpass",
    )
    headers = auth_headers("manager@test.local", "managerpass")

    allowed_response = client.post(
        "/users/",
        json={
            "full_name": "Employee User",
            "email": "employee-created@test.local",
            "password": "password123",
            "role_id": 3,
            "is_active": True,
        },
        headers=headers,
    )
    assert allowed_response.status_code == 200
    assert allowed_response.json()["role_id"] == 3

    forbidden_response = client.post(
        "/users/",
        json={
            "full_name": "Admin User",
            "email": "admin-created@test.local",
            "password": "password123",
            "role_id": 1,
            "is_active": True,
        },
        headers=headers,
    )
    assert forbidden_response.status_code == 403


def test_admin_can_activate_imported_user_with_one_time_temporary_password(
    client,
    auth_headers,
    factories,
    db_session,
    app_modules,
):
    factories["create_user"](
        full_name="Admin User",
        email="admin@test.local",
        role_id=1,
        password="adminpass",
    )
    factories["create_user"](
        full_name="Manager User",
        email="manager@test.local",
        role_id=2,
        password="managerpass",
    )
    imported_user = factories["create_user"](
        full_name="Imported User",
        email="imported@test.local",
        role_id=3,
        password="hidden-random-password",
        is_active=False,
    )
    imported_user.must_change_password = True
    db_session.commit()

    manager_response = client.post(
        f"/users/{imported_user.id}/reset-password",
        headers=auth_headers("manager@test.local", "managerpass"),
    )
    assert manager_response.status_code == 403

    admin_response = client.post(
        f"/users/{imported_user.id}/reset-password",
        headers=auth_headers("admin@test.local", "adminpass"),
    )
    assert admin_response.status_code == 200
    data = admin_response.json()
    assert data["temporary_password"]
    assert data["user"]["is_active"] is True
    assert data["user"]["must_change_password"] is True

    db_session.refresh(imported_user)
    assert imported_user.is_active is True
    assert imported_user.must_change_password is True
    assert app_modules["auth"].verify_password(
        data["temporary_password"],
        imported_user.password_hash,
    )


def test_change_password_requires_current_password_and_clears_flag(
    client,
    auth_headers,
    factories,
    db_session,
):
    user = factories["create_user"](
        full_name="Password User",
        email="password@test.local",
        role_id=3,
        password="oldpassword",
    )
    user.must_change_password = True
    db_session.commit()
    headers = auth_headers("password@test.local", "oldpassword")

    wrong_current_response = client.post(
        "/auth/change-password",
        json={"current_password": "badpassword", "new_password": "newpassword123"},
        headers=headers,
    )
    assert wrong_current_response.status_code == 400

    same_password_response = client.post(
        "/auth/change-password",
        json={"current_password": "oldpassword", "new_password": "oldpassword"},
        headers=headers,
    )
    assert same_password_response.status_code == 400

    short_password_response = client.post(
        "/auth/change-password",
        json={"current_password": "oldpassword", "new_password": "short"},
        headers=headers,
    )
    assert short_password_response.status_code == 422

    success_response = client.post(
        "/auth/change-password",
        json={"current_password": "oldpassword", "new_password": "newpassword123"},
        headers=headers,
    )
    assert success_response.status_code == 200

    db_session.refresh(user)
    assert user.must_change_password is False

    login_response = client.post(
        "/auth/login",
        data={"username": "password@test.local", "password": "newpassword123"},
    )
    assert login_response.status_code == 200
    assert login_response.json()["user"]["must_change_password"] is False


def test_employee_cannot_access_admin_client_endpoints(client, auth_headers, factories):
    factories["create_user"](
        full_name="Employee User",
        email="employee@test.local",
        role_id=3,
        password="employeepass",
    )

    response = client.post(
        "/clients/",
        json={"company_name": "Forbidden Client"},
        headers=auth_headers("employee@test.local", "employeepass"),
    )
    assert response.status_code == 403


def test_unauthenticated_protected_endpoints_return_401(client):
    protected_requests = [
        ("get", "/auth/me", None),
        ("get", "/tasks/", None),
        ("get", "/users/", None),
        ("get", "/logs/", None),
        ("get", "/dashboard/stats", None),
        ("get", "/excel/global/template", None),
    ]

    for method, path, payload in protected_requests:
        request = getattr(client, method)
        response = request(path, json=payload) if payload is not None else request(path)
        assert response.status_code == 401, f"{method.upper()} {path} returned {response.status_code}"


def test_employee_cannot_create_or_delete_tasks(client, auth_headers, factories):
    admin = factories["create_user"](
        full_name="Admin RBAC",
        email="admin-rbac@test.local",
        role_id=1,
        password="adminpass",
    )
    employee = factories["create_user"](
        full_name="Employee RBAC",
        email="employee-rbac@test.local",
        role_id=3,
        password="employeepass",
    )
    client_record = factories["create_client"](company_name="RBAC Client")
    task = factories["create_task"](
        title="RBAC Task",
        client_id=client_record.id,
        created_by=admin.id,
        assigned_to=employee.id,
    )
    headers = auth_headers("employee-rbac@test.local", "employeepass")

    create_response = client.post(
        "/tasks/",
        json={"title": "Forbidden Task", "client_id": str(client_record.id)},
        headers=headers,
    )
    assert create_response.status_code == 403

    delete_response = client.delete(f"/tasks/{task.id}", headers=headers)
    assert delete_response.status_code == 403


def test_employee_cannot_access_logs_or_user_management(client, auth_headers, factories):
    factories["create_user"](
        full_name="Employee Restricted",
        email="employee-restricted@test.local",
        role_id=3,
        password="employeepass",
    )
    headers = auth_headers("employee-restricted@test.local", "employeepass")

    logs_response = client.get("/logs/", headers=headers)
    assert logs_response.status_code == 403

    users_list_response = client.get("/users/", headers=headers)
    assert users_list_response.status_code == 403

    users_create_response = client.post(
        "/users/",
        json={
            "full_name": "Forbidden User",
            "email": "forbidden-user@test.local",
            "password": "password123",
            "role_id": 3,
            "is_active": True,
        },
        headers=headers,
    )
    assert users_create_response.status_code == 403


def test_manager_cannot_reset_password_or_import_excel(client, auth_headers, factories):
    factories["create_user"](
        full_name="Manager RBAC",
        email="manager-rbac@test.local",
        role_id=2,
        password="managerpass",
    )
    imported_user = factories["create_user"](
        full_name="Imported RBAC",
        email="imported-rbac@test.local",
        role_id=3,
        password="hiddenpassword",
        is_active=False,
    )
    headers = auth_headers("manager-rbac@test.local", "managerpass")

    reset_response = client.post(f"/users/{imported_user.id}/reset-password", headers=headers)
    assert reset_response.status_code == 403

    workbook = _make_global_workbook_bytes(
        ["Tapşırıq", "İcraçı", "Prioritet", "Status"],
        [["Manager Import Task", "Imported RBAC", "Orta", "Yeni"]],
    )
    import_response = client.post(
        "/excel/global/import",
        files={
            "file": (
                "manager-import.xlsx",
                workbook,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )
    assert import_response.status_code == 403


def test_admin_can_use_restricted_management_endpoints(client, auth_headers, factories):
    admin = factories["create_user"](
        full_name="Admin Allowed",
        email="admin-allowed@test.local",
        role_id=1,
        password="adminpass",
    )
    client_record = factories["create_client"](company_name="Admin Allowed Client")
    headers = auth_headers("admin-allowed@test.local", "adminpass")

    task_create_response = client.post(
        "/tasks/",
        json={"title": "Admin Created Task", "client_id": str(client_record.id)},
        headers=headers,
    )
    assert task_create_response.status_code == 200

    logs_response = client.get("/logs/", headers=headers)
    assert logs_response.status_code == 200

    user_create_response = client.post(
        "/users/",
        json={
            "full_name": "Admin Created Employee",
            "email": "admin-created-employee@test.local",
            "password": "password123",
            "role_id": 3,
            "is_active": True,
        },
        headers=headers,
    )
    assert user_create_response.status_code == 200

    import_workbook = _make_global_workbook_bytes(
        ["Tapşırıq", "İcraçı", "Prioritet", "Status"],
        [["Admin Import Task", "Admin Created Employee", "Orta", "Yeni"]],
    )
    import_response = client.post(
        "/excel/global/import",
        files={
            "file": (
                "admin-import.xlsx",
                import_workbook,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )
    assert import_response.status_code == 200

    delete_response = client.delete(
        f"/tasks/{task_create_response.json()['id']}",
        headers=auth_headers("admin-allowed@test.local", "adminpass"),
    )
    assert delete_response.status_code == 200

    assert admin.role_id == 1


def test_client_create_and_read(client, auth_headers, factories):
    factories["create_user"](
        full_name="Admin User",
        email="admin@test.local",
        role_id=1,
        password="adminpass",
    )
    headers = auth_headers("admin@test.local", "adminpass")

    create_response = client.post(
        "/clients/",
        json={
            "company_name": "Acme LLC",
            "contact_person": "Nigar",
            "email": "contact@acme.test",
        },
        headers=headers,
    )
    assert create_response.status_code == 200
    client_id = create_response.json()["id"]

    list_response = client.get("/clients/", headers=headers)
    assert list_response.status_code == 200
    assert len(list_response.json()) == 1

    detail_response = client.get(f"/clients/{client_id}", headers=headers)
    assert detail_response.status_code == 200
    assert detail_response.json()["company_name"] == "Acme LLC"


def test_task_create_read_and_employee_status_update(
    client,
    auth_headers,
    factories,
    app_modules,
    db_session,
):
    admin = factories["create_user"](
        full_name="Admin User",
        email="admin@test.local",
        role_id=1,
        password="adminpass",
    )
    employee = factories["create_user"](
        full_name="Employee User",
        email="employee@test.local",
        role_id=3,
        password="employeepass",
    )
    client_record = factories["create_client"](company_name="Task Client")

    create_response = client.post(
        "/tasks/",
        json={
            "title": "Tax report",
            "client_id": str(client_record.id),
            "assigned_to": str(employee.id),
            "status": "todo",
            "priority": "high",
            "start_date": "2026-01-05",
        },
        headers=auth_headers("admin@test.local", "adminpass"),
    )
    assert create_response.status_code == 200
    task_id = create_response.json()["id"]
    assert create_response.json()["created_by"] == str(admin.id)
    assert create_response.json()["start_date"] == "2026-01-05"
    models = app_modules["models"]
    created_task = db_session.query(models.Task).filter(models.Task.id == task_id).first()
    assert created_task.sxa_task_id is None
    assert created_task.start_date.isoformat() == "2026-01-05"

    list_response = client.get("/tasks/", headers=auth_headers("admin@test.local", "adminpass"))
    assert list_response.status_code == 200
    assert len(list_response.json()) == 1

    update_response = client.put(
        f"/tasks/{task_id}",
        json={"status": "done"},
        headers=auth_headers("employee@test.local", "employeepass"),
    )
    assert update_response.status_code == 200
    assert update_response.json()["status"] == "done"


def test_attachment_upload_sanitizes_filename_and_downloads(
    client,
    auth_headers,
    factories,
    db_session,
    app_modules,
):
    admin = factories["create_user"](
        full_name="Admin Upload",
        email="admin-upload@test.local",
        role_id=1,
        password="adminpass",
    )
    client_record = factories["create_client"](company_name="Upload Client")
    task = factories["create_task"](
        title="Upload Task",
        client_id=client_record.id,
        created_by=admin.id,
    )
    headers = auth_headers("admin-upload@test.local", "adminpass")

    upload_response = client.post(
        f"/tasks/{task.id}/attachments",
        files={
            "file": (
                "../unsafe report.txt",
                BytesIO(b"plain text attachment"),
                "text/plain",
            )
        },
        headers=headers,
    )
    assert upload_response.status_code == 200
    data = upload_response.json()
    assert data["file_name"] == "unsafe_report.txt"
    assert ".." not in data["file_path"]

    models = app_modules["models"]
    attachment = db_session.query(models.TaskAttachment).filter(
        models.TaskAttachment.id == data["id"]
    ).first()
    assert attachment is not None
    assert attachment.file_path.endswith("unsafe_report.txt")

    download_response = client.get(
        f"/attachments/{data['id']}/download",
        headers=headers,
    )
    assert download_response.status_code == 200
    assert download_response.content == b"plain text attachment"


def test_attachment_upload_rejects_mismatched_content_signature(
    client,
    auth_headers,
    factories,
):
    admin = factories["create_user"](
        full_name="Admin Fake PDF",
        email="admin-fake-pdf@test.local",
        role_id=1,
        password="adminpass",
    )
    client_record = factories["create_client"](company_name="Fake PDF Client")
    task = factories["create_task"](
        title="Fake PDF Task",
        client_id=client_record.id,
        created_by=admin.id,
    )

    response = client.post(
        f"/tasks/{task.id}/attachments",
        files={"file": ("fake.pdf", BytesIO(b"not a real pdf"), "application/pdf")},
        headers=auth_headers("admin-fake-pdf@test.local", "adminpass"),
    )
    assert response.status_code == 400


def test_attachment_download_blocks_paths_outside_upload_root(
    client,
    auth_headers,
    factories,
    db_session,
    app_modules,
    tmp_path,
):
    admin = factories["create_user"](
        full_name="Admin Path",
        email="admin-path@test.local",
        role_id=1,
        password="adminpass",
    )
    client_record = factories["create_client"](company_name="Path Client")
    task = factories["create_task"](
        title="Path Task",
        client_id=client_record.id,
        created_by=admin.id,
    )
    outside_file = tmp_path / "outside.txt"
    outside_file.write_text("outside", encoding="utf-8")

    models = app_modules["models"]
    attachment = models.TaskAttachment(
        task_id=task.id,
        uploaded_by=admin.id,
        file_name="outside.txt",
        file_path=str(outside_file),
        file_type="text/plain",
    )
    db_session.add(attachment)
    db_session.commit()
    db_session.refresh(attachment)

    response = client.get(
        f"/attachments/{attachment.id}/download",
        headers=auth_headers("admin-path@test.local", "adminpass"),
    )
    assert response.status_code == 400


# ---------------------------------------------------------------------------
# Phase 8E: category field tests
# ---------------------------------------------------------------------------


def test_task_create_and_read_with_category(client, auth_headers, factories, db_session, app_modules):
    """Task create stores category; response includes category."""
    factories["create_user"](
        full_name="Admin Category",
        email="admin-cat@test.local",
        role_id=1,
        password="adminpass",
    )
    factories["create_user"](
        full_name="Worker Cat",
        email="worker-cat@test.local",
        role_id=3,
    )
    client_record = factories["create_client"](company_name="Category Client")

    models = app_modules["models"]
    worker = db_session.query(models.User).filter(models.User.email == "worker-cat@test.local").first()

    # Create with category
    create_response = client.post(
        "/tasks/",
        json={
            "title": "Category Task",
            "category": "Maliyyə",
            "client_id": str(client_record.id),
            "assigned_to": str(worker.id),
            "priority": "medium",
        },
        headers=auth_headers("admin-cat@test.local", "adminpass"),
    )
    assert create_response.status_code == 200
    data = create_response.json()
    assert data["category"] == "Maliyyə"
    task_id = data["id"]

    # Verify DB
    created = db_session.query(models.Task).filter(models.Task.id == task_id).first()
    assert created.category == "Maliyyə"
    assert created.direction_program is None  # New writes only go to category, not legacy field

    # Update category
    update_response = client.put(
        f"/tasks/{task_id}",
        json={"category": "İT"},
        headers=auth_headers("admin-cat@test.local", "adminpass"),
    )
    assert update_response.status_code == 200
    assert update_response.json()["category"] == "İT"
    db_session.refresh(created)
    assert created.category == "İT"


def test_task_create_with_direction_program_maps_to_category(
    client, auth_headers, factories, db_session, app_modules,
):
    """Old clients sending direction_program get category populated."""
    factories["create_user"](
        full_name="Admin Backward",
        email="admin-bw@test.local",
        role_id=1,
        password="adminpass",
    )
    client_record = factories["create_client"](company_name="BW Client")

    models = app_modules["models"]

    create_response = client.post(
        "/tasks/",
        json={
            "title": "Backward Compat Task",
            "direction_program": "LegacyCat",
            "client_id": str(client_record.id),
        },
        headers=auth_headers("admin-bw@test.local", "adminpass"),
    )
    assert create_response.status_code == 200
    data = create_response.json()
    assert data["category"] == "LegacyCat"

    created = (
        db_session.query(models.Task)
        .filter(models.Task.id == data["id"])
        .first()
    )
    assert created.category == "LegacyCat"
    assert created.direction_program is None  # Phase 8E.1: direction_program is now read-only legacy fallback.


# ---------------------------------------------------------------------------
# Phase 6A: custom_fields tests
# ---------------------------------------------------------------------------


def test_task_custom_fields_create_without_custom_fields(
    client, auth_headers, factories,
):
    """Create task without custom_fields → succeeds, returns {}."""
    factories["create_user"](
        full_name="Admin CF", email="admin-cf1@test.local", role_id=1, password="adminpass"
    )
    headers = auth_headers("admin-cf1@test.local", "adminpass")

    resp = client.post(
        "/tasks/",
        json={"title": "No Custom Fields Task"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["title"] == "No Custom Fields Task"
    # custom_fields should default to empty dict (or null → frontend treats as {})
    assert data.get("custom_fields") in (None, {})
    if data.get("custom_fields") is not None:
        assert data["custom_fields"] == {}


def test_task_custom_fields_create_with_custom_fields(
    client, auth_headers, factories, db_session, app_modules,
):
    """Create task with custom_fields → saves and returns values."""
    factories["create_user"](
        full_name="Admin CF", email="admin-cf2@test.local", role_id=1, password="adminpass"
    )
    headers = auth_headers("admin-cf2@test.local", "adminpass")

    custom = {"branch": "Bakı", "contract_no": "CN-102", "payment_status": "Ödənilib"}
    resp = client.post(
        "/tasks/",
        json={"title": "Custom Fields Task", "custom_fields": custom},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["title"] == "Custom Fields Task"
    assert data["custom_fields"] == custom

    # Verify in DB
    models = app_modules["models"]
    task = (
        db_session.query(models.Task)
        .filter(models.Task.title == "Custom Fields Task")
        .first()
    )
    assert task is not None
    assert task.custom_fields == custom


def test_task_custom_fields_update(
    client, auth_headers, factories, db_session, app_modules,
):
    """Update task custom_fields → persists changes."""
    factories["create_user"](
        full_name="Admin CF", email="admin-cf3@test.local", role_id=1, password="adminpass"
    )
    headers = auth_headers("admin-cf3@test.local", "adminpass")

    # Create with initial custom_fields
    resp = client.post(
        "/tasks/",
        json={"title": "Update CF Task", "custom_fields": {"key1": "value1"}},
        headers=headers,
    )
    assert resp.status_code == 200
    task_id = resp.json()["id"]

    # Update custom_fields
    resp2 = client.put(
        f"/tasks/{task_id}",
        json={"custom_fields": {"key1": "updated", "key2": "new_value"}},
        headers=headers,
    )
    assert resp2.status_code == 200, resp2.text
    assert resp2.json()["custom_fields"] == {"key1": "updated", "key2": "new_value"}

    # Clear custom_fields
    resp3 = client.put(
        f"/tasks/{task_id}",
        json={"custom_fields": {}},
        headers=headers,
    )
    assert resp3.status_code == 200, resp3.text
    assert resp3.json()["custom_fields"] == {}

    # Verify in DB
    models = app_modules["models"]
    task = db_session.query(models.Task).filter(models.Task.id == task_id).first()
    assert task.custom_fields == {}


def test_task_custom_fields_rejects_non_dict(
    client, auth_headers, factories,
):
    """custom_fields as list/string/number → validation error."""
    factories["create_user"](
        full_name="Admin CF", email="admin-cf4@test.local", role_id=1, password="adminpass"
    )
    headers = auth_headers("admin-cf4@test.local", "adminpass")

    for bad_value in (["should be dict"], "string", 123, True):
        resp = client.post(
            "/tasks/",
            json={"title": f"Bad CF {type(bad_value).__name__}", "custom_fields": bad_value},
            headers=headers,
        )
        assert resp.status_code == 422, (
            f"Expected 422 for {type(bad_value).__name__}, got {resp.status_code}: {resp.text}"
        )


def test_task_list_detail_works_with_custom_fields(
    client, auth_headers, factories,
):
    """Task list and detail endpoints return custom_fields."""
    factories["create_user"](
        full_name="Admin CF", email="admin-cf5@test.local", role_id=1, password="adminpass"
    )
    headers = auth_headers("admin-cf5@test.local", "adminpass")

    # Create task with custom_fields
    resp = client.post(
        "/tasks/",
        json={"title": "List CF Task", "custom_fields": {"dept": "IT"}},
        headers=headers,
    )
    assert resp.status_code == 200
    task_id = resp.json()["id"]

    # List
    list_resp = client.get("/tasks/", headers=headers)
    assert list_resp.status_code == 200
    tasks = list_resp.json()
    cf_task = next((t for t in tasks if t["id"] == task_id), None)
    assert cf_task is not None
    assert cf_task["custom_fields"] == {"dept": "IT"}

    # Detail
    detail_resp = client.get(f"/tasks/{task_id}", headers=headers)
    assert detail_resp.status_code == 200
    assert detail_resp.json()["custom_fields"] == {"dept": "IT"}


def test_employee_only_sees_own_tasks(client, auth_headers, factories):
    admin = factories["create_user"](
        full_name="Admin User",
        email="admin@test.local",
        role_id=1,
        password="adminpass",
    )
    employee_one = factories["create_user"](
        full_name="Employee One",
        email="employee1@test.local",
        role_id=3,
        password="password123",
    )
    employee_two = factories["create_user"](
        full_name="Employee Two",
        email="employee2@test.local",
        role_id=3,
        password="password123",
    )
    client_record = factories["create_client"](company_name="Scoped Client")
    factories["create_task"](
        title="Task for employee one",
        client_id=client_record.id,
        created_by=admin.id,
        assigned_to=employee_one.id,
    )
    factories["create_task"](
        title="Task for employee two",
        client_id=client_record.id,
        created_by=admin.id,
        assigned_to=employee_two.id,
    )

    response = client.get(
        "/tasks/",
        headers=auth_headers("employee1@test.local", "password123"),
    )
    assert response.status_code == 200
    assert len(response.json()) == 1
    assert response.json()[0]["title"] == "Task for employee one"


def test_excel_user_helper_creates_deterministic_email_once(db_session, app_modules):
    from app.routers.excel import get_or_create_user_from_excel

    models = app_modules["models"]

    user, created = get_or_create_user_from_excel(db_session, "  Turan   Məmmədov  ")
    assert created is True
    assert user.full_name == "Turan Məmmədov"
    assert user.email == "turan.mammadov@task-manager.local"
    assert user.role_id == 3
    assert user.is_active is False
    assert user.must_change_password is True

    reused_user, reused_created = get_or_create_user_from_excel(db_session, "Turan Məmmədov")
    assert reused_created is False
    assert reused_user.id == user.id
    assert db_session.query(models.User).filter(models.User.email == "turan.mammadov@task-manager.local").count() == 1


def test_excel_user_helper_reuses_case_insensitive_and_legacy_email(db_session, factories):
    from app.routers.excel import get_or_create_user_from_excel

    existing_user = factories["create_user"](
        full_name="Turan Məmmədov",
        email="TURAN.MAMMADOV@TASK-MANAGER.LOCAL",
        role_id=3,
    )
    reused_user, created = get_or_create_user_from_excel(db_session, "Turan Məmmədov")
    assert created is False
    assert reused_user.id == existing_user.id

    legacy_user = factories["create_user"](
        full_name="Samir Məmmədov",
        email="samir.memmedov@task-manager.local",
        role_id=3,
    )
    reused_legacy_user, legacy_created = get_or_create_user_from_excel(db_session, "Samir Məmmədov")
    assert legacy_created is False
    assert reused_legacy_user.id == legacy_user.id


def test_excel_export_uses_neutral_import_workbook_shape(client, auth_headers, factories, db_session, app_modules):
    from datetime import date, datetime

    models = app_modules["models"]
    admin = factories["create_user"](
        full_name="Admin User",
        email="admin-export@test.local",
        role_id=1,
        password="adminpass",
    )
    assignee = factories["create_user"](
        full_name="Worker User",
        email="worker-export@test.local",
        role_id=3,
    )
    project = models.Project(
        code="PRJ-1",
        name="Export Project",
        direction_program="Audit",
        owner="Owner User",
        approach="Agile",
    )
    alpha_project = models.Project(
        code="PRJ-0",
        name="Alpha Project",
        direction_program="Strategy",
        owner="Alpha Owner",
        approach="Waterfall",
    )
    db_session.add_all([project, alpha_project])
    db_session.flush()
    stage = models.ProjectStage(
        project_id=project.id,
        name="İcra",
        weight_percent=50,
        planned_start=date(2026, 1, 1),
        planned_end=date(2026, 1, 31),
        actual_end=None,
        status="in_progress",
        delay_reason="",
        completion_percent=25,
        sort_order=3,
    )
    alpha_stage_late = models.ProjectStage(
        project_id=alpha_project.id,
        name="Bağlanış",
        weight_percent=20,
        planned_start=date(2026, 3, 1),
        planned_end=date(2026, 3, 15),
        status="completed",
        sort_order=5,
    )
    alpha_stage_early = models.ProjectStage(
        project_id=alpha_project.id,
        name="Planlaşdırma",
        weight_percent=30,
        planned_start=date(2026, 2, 1),
        planned_end=date(2026, 2, 15),
        status="not_started",
        sort_order=2,
    )
    long_title = "Export Task " + ("long description " * 8)
    task = models.Task(
        title=long_title,
        description="Export description",
        project_id=project.id,
        created_by=admin.id,
        assigned_to=assignee.id,
        task_type="recurring",
        operational_status="active",
        status="in_progress",
        priority="high",
        kanban_enabled=True,
        frequency="monthly",
        reporting_required=False,
        category="Audit",
        direction_program="Audit",
        risk_status="medium",
        external_id="EXT-EXPORT-1",
        sxa_task_id="SXA-1",
        start_date=date(2026, 1, 4),
        sxa_start_date=date(2026, 1, 5),
        initial_deadline=date(2026, 1, 20),
        current_deadline=date(2026, 1, 25),
        deadline=datetime(2026, 1, 25),
        custom_fields={"Filial": "Bakı", "Müqavilə nömrəsi": "CN-102"},
    )
    legacy_id_task = models.Task(
        title="Task with legacy export fallback",
        project_id=project.id,
        created_by=admin.id,
        assigned_to=assignee.id,
        operational_status="waiting",
        status="todo",
        priority="medium",
        sxa_task_id="SXA-FALLBACK-EXPORT",
        sxa_start_date=date(2026, 2, 1),
    )
    db_session.add_all([stage, alpha_stage_late, alpha_stage_early, task, legacy_id_task])
    db_session.commit()

    response = client.get(
        "/excel/tasks/export",
        headers=auth_headers("admin-export@test.local", "adminpass"),
    )

    assert response.status_code == 200
    assert "task-manager-export-" in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content), data_only=True)
    try:
        assert workbook.sheetnames == [
            "Tapşırıqlar",
            "Qruplar",
            "Mərhələlər",
        ]
        master_rows = list(workbook["Tapşırıqlar"].iter_rows(values_only=True))
        project_rows = list(workbook["Qruplar"].iter_rows(values_only=True))
        stage_rows = list(workbook["Mərhələlər"].iter_rows(values_only=True))

        assert master_rows[0][:12] == (
            "Xarici ID",
            "Tapşırıq",
            "Təsvir",
            "İcraçı",
            "Prioritet",
            "Status",
            "Başlama tarixi",
            "Son tarix",
            "Kateqoriya",
            "Kanban",
            "Filial",
            "Müqavilə nömrəsi",
        )
        assert master_rows[1][:12] == (
            "EXT-EXPORT-1",
            long_title,
            "Export description",
            "Worker User",
            "Yüksək",
            "İcradadır",
            datetime(2026, 1, 4),
            datetime(2026, 1, 25),
            "Audit",
            "Bəli",
            "Bakı",
            "CN-102",
        )
        assert master_rows[2][0] == "SXA-FALLBACK-EXPORT"
        assert master_rows[2][1] == "Task with legacy export fallback"
        assert master_rows[2][6].date() == date(2026, 2, 1)
        assert ("PRJ-1", "Export Project", "Audit", "Owner User", "Agile") in project_rows
        assert stage_rows[1][:2] == ("Alpha Project", "Planlaşdırma")
        assert stage_rows[2][:2] == ("Alpha Project", "Bağlanış")
        assert stage_rows[3][3].date() == date(2026, 1, 1)
        assert stage_rows[3][4].date() == date(2026, 1, 31)
        assert stage_rows[3][:3] == (
            "Export Project",
            "İcra",
            50,
        )
        assert stage_rows[3][5:7] == (None, "İcradadır")
        assert workbook["Tapşırıqlar"].freeze_panes == "A2"
        assert workbook["Tapşırıqlar"].cell(row=1, column=1).fill.fgColor.rgb.endswith("1F4E78")
        assert workbook["Tapşırıqlar"].cell(row=2, column=6).fill.fgColor.rgb.endswith("DCFCE7")
        assert workbook["Tapşırıqlar"].row_dimensions[2].height > 22
        assert workbook["Tapşırıqlar"].column_dimensions["B"].width > 20
    finally:
        workbook.close()


def test_ai_readonly_tool_does_not_require_confirmation(
    client,
    auth_headers,
    factories,
    monkeypatch,
):
    admin = factories["create_user"](
        full_name="Admin User",
        email="admin@test.local",
        role_id=1,
        password="adminpass",
    )
    client_record = factories["create_client"](company_name="AI Client")
    factories["create_task"](
        title="AI visible task",
        client_id=client_record.id,
        created_by=admin.id,
        assigned_to=admin.id,
    )

    from app.ai import agent

    fake_client = _FakeOpenAIClient(
        [
            _make_function_response(
                response_id="resp_1",
                tool_name="read_tasks",
                arguments={"query": "AI", "limit": 10},
            ),
            _make_text_response(
                response_id="resp_2",
                text="1 tapşırıq tapıldı.",
            ),
        ]
    )
    monkeypatch.setattr(agent, "get_openai_client", lambda: fake_client)

    response = client.post(
        "/ai/assistant-v2/message",
        json={"message": "AI tasklarını göstər"},
        headers=auth_headers("admin@test.local", "adminpass"),
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "completed"
    assert payload["tool_results"][0]["tool"] == "read_tasks"
    assert payload.get("requires_confirmation") in (None, False)


def test_dashboard_summary_uses_business_status_for_waiting_and_delay_metrics(
    db_session,
    factories,
):
    from datetime import date, timedelta

    from app.ai.tools import get_dashboard_summary

    admin = factories["create_user"](
        full_name="Admin User",
        email="summary-admin@test.local",
        role_id=1,
    )
    client_record = factories["create_client"](company_name="Summary Client")
    today = date.today()

    def make_task(title, *, operational_status, current_deadline=None, status="todo"):
        task = factories["create_task"](
            title=title,
            client_id=client_record.id,
            created_by=admin.id,
            assigned_to=admin.id,
            status=status,
        )
        task.operational_status = operational_status
        task.current_deadline = current_deadline
        db_session.commit()
        db_session.refresh(task)
        return task

    make_task("Business waiting", operational_status="waiting")
    make_task("Technical todo only", operational_status=None)
    make_task("Within deadline", operational_status="within_deadline")
    make_task(
        "Official delayed",
        operational_status="overdue",
        current_deadline=today + timedelta(days=2),
    )
    make_task(
        "Date overdue not marked",
        operational_status="active",
        current_deadline=today - timedelta(days=1),
    )
    make_task(
        "Official and date overdue",
        operational_status="overdue",
        current_deadline=today - timedelta(days=2),
    )
    make_task(
        "Completed past deadline",
        operational_status="completed",
        current_deadline=today - timedelta(days=3),
        status="done",
    )

    result = get_dashboard_summary(db_session, admin)
    summary = result["summary"]

    assert summary["gozlemede"] == 1
    assert summary["todo_tasks"] == 1
    assert summary["official_delayed_tasks"] == 2
    assert summary["date_overdue_tasks"] == 2
    assert summary["date_overdue_not_marked_tasks"] == 1
    assert summary["total_delayed_tasks"] == 3
    assert summary["overdue_tasks"] == 3
    assert summary["status_bolgusu"]["gecikme_statusunda"] == 2


def test_ai_create_task_requires_confirmation_before_execution(
    client,
    auth_headers,
    factories,
    db_session,
    monkeypatch,
):
    admin = factories["create_user"](
        full_name="Admin User",
        email="admin@test.local",
        role_id=1,
        password="adminpass",
    )
    assignee = factories["create_user"](
        full_name="Worker User",
        email="worker@test.local",
        role_id=3,
        password="workerpass",
    )
    client_record = factories["create_client"](company_name="Confirm Client")

    from app.ai import agent
    from app import models

    fake_client = _FakeOpenAIClient(
        [
            _make_function_response(
                response_id="resp_1",
                tool_name="create_task",
                arguments={
                    "title": "Yeni task",
                    "client_id": str(client_record.id),
                    "assigned_to": str(assignee.id),
                    "priority": "high",
                },
            )
        ]
    )
    monkeypatch.setattr(agent, "get_openai_client", lambda: fake_client)

    response = client.post(
        "/ai/assistant-v2/message",
        json={"message": "Yeni task yarat"},
        headers=auth_headers("admin@test.local", "adminpass"),
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "pending_confirmation"
    assert payload["pending_tool"]["tool_name"] == "create_task"
    assert db_session.query(models.Task).count() == 0


def test_ai_confirm_token_executes_create_task(
    client,
    auth_headers,
    factories,
    db_session,
    monkeypatch,
):
    admin = factories["create_user"](
        full_name="Admin User",
        email="admin@test.local",
        role_id=1,
        password="adminpass",
    )
    assignee = factories["create_user"](
        full_name="Worker User",
        email="worker@test.local",
        role_id=3,
        password="workerpass",
    )
    client_record = factories["create_client"](company_name="Confirm Client")

    from app.ai import agent
    from app import models

    fake_client = _FakeOpenAIClient(
        [
            _make_function_response(
                response_id="resp_1",
                tool_name="create_task",
                arguments={
                    "title": "Təsdiqli task",
                    "client_id": str(client_record.id),
                    "assigned_to": str(assignee.id),
                    "priority": "medium",
                },
            )
        ]
    )
    monkeypatch.setattr(agent, "get_openai_client", lambda: fake_client)

    initial_response = client.post(
        "/ai/assistant-v2/message",
        json={"message": "Task yarat"},
        headers=auth_headers("admin@test.local", "adminpass"),
    )
    assert initial_response.status_code == 200
    confirmation_token = initial_response.json()["confirmation_token"]

    confirm_response = client.post(
        "/ai/assistant-v2/confirm",
        json={"confirmation_token": confirmation_token},
        headers=auth_headers("admin@test.local", "adminpass"),
    )

    assert confirm_response.status_code == 200
    payload = confirm_response.json()
    assert payload["status"] == "completed"
    assert payload["tool_results"][0]["tool"] == "create_task"

    created_task = db_session.query(models.Task).one()
    assert created_task.title == "Təsdiqli task"
    assert created_task.assigned_to == assignee.id


# ──────────────────────────────────────────────────────────────────────
# Resolver tests (centralized resolver layer from app.ai.resolvers)
# ──────────────────────────────────────────────────────────────────────


def test_resolve_priority_azerbaijani():
    """resolve_priority maps Azerbaijani labels correctly."""
    from app.ai.resolvers import resolve_priority

    assert resolve_priority("aşağı") == "low"
    assert resolve_priority("asagi") == "low"
    assert resolve_priority("zəif") == "low"
    assert resolve_priority("yüksək") == "high"
    assert resolve_priority("yuksek") == "high"
    assert resolve_priority("təcili") == "high"
    assert resolve_priority("vacib") == "high"
    assert resolve_priority("orta") == "medium"
    assert resolve_priority("medium") == "medium"
    assert resolve_priority("low") == "low"
    assert resolve_priority("high") == "high"
    assert resolve_priority(None) is None
    assert resolve_priority("") is None
    assert resolve_priority("naməlum") is None


def test_resolve_task_type_azerbaijani():
    """resolve_task_type maps Azerbaijani labels correctly."""
    from app.ai.resolvers import resolve_task_type

    assert resolve_task_type("birdəfəlik") == "one_time"
    assert resolve_task_type("birdefelik") == "one_time"
    assert resolve_task_type("one_time") == "one_time"
    assert resolve_task_type("one time") == "one_time"
    assert resolve_task_type("ümumi") == "general"
    assert resolve_task_type("general") == "general"
    assert resolve_task_type("təkrarlanan") == "recurring"
    assert resolve_task_type("recurring") == "recurring"
    assert resolve_task_type("layihə") == "project"
    assert resolve_task_type("project") == "project"
    assert resolve_task_type(None) is None
    assert resolve_task_type("") is None


def test_resolve_operational_status():
    """resolve_operational_status maps Azerbaijani labels correctly."""
    from app.ai.resolvers import resolve_operational_status

    assert resolve_operational_status("yeni") == "new"
    assert resolve_operational_status("gözləmədə") == "waiting"
    assert resolve_operational_status("gözlemede") == "waiting"
    assert resolve_operational_status("icra müddətində") == "within_deadline"
    assert resolve_operational_status("icradadır") == "active"
    assert resolve_operational_status("aktiv") == "active"
    assert resolve_operational_status("gecikmə") == "overdue"
    assert resolve_operational_status("gecikib") == "overdue"
    assert resolve_operational_status("tam icra") == "completed"
    assert resolve_operational_status("tamamlanıb") == "completed"
    # Kanban aliases
    assert resolve_operational_status("todo") == "waiting"
    assert resolve_operational_status("in_progress") == "active"
    assert resolve_operational_status("done") == "completed"


def test_normalize_az_text():
    """normalize_az_text folds Azerbaijani characters and cleans whitespace."""
    from app.ai.resolvers import normalize_az_text

    assert normalize_az_text("Nərminə") == "nermine"
    assert normalize_az_text("  Hacı  YAĞIZ  ") == "haci yagiz"
    assert normalize_az_text("L026") == "l026"
    assert normalize_az_text("") == ""
    assert normalize_az_text(None) == ""
    # Punctuation removed, codes kept
    result = normalize_az_text("Şahin! test.")
    assert "sahin" in result
    assert "test" in result


def test_strip_az_suffixes():
    """strip_az_suffixes removes Azerbaijani dative suffixes from names."""
    from app.ai.resolvers import strip_az_suffixes

    assert strip_az_suffixes("Nərminə") == "Nərmin"
    assert strip_az_suffixes("Elminə") == "Elmin"
    assert strip_az_suffixes("Fəridə") == "Fərid"
    assert strip_az_suffixes("Hacıya") == "Hacı"
    assert strip_az_suffixes("Hacığa") == "Hacı"
    assert strip_az_suffixes("Əliyə") == "Əli"
    # Name without suffix passes through
    assert strip_az_suffixes("Turan") == "Turan"
    # Very short names should not be over-stripped
    stripped = strip_az_suffixes("Eli")
    assert len(stripped) >= 2
    assert strip_az_suffixes("") == ""
    assert strip_az_suffixes(None) == ""


def test_fuzzy_match_exact_and_normalized():
    """fuzzy_match returns single / ambiguous / none correctly."""
    from app.ai.resolvers import fuzzy_match

    candidates = ["Nərmin Əmirova", "Nərgiz Əliyeva", "Turan Məmmədov"]

    # Exact raw match
    result = fuzzy_match("Nərmin Əmirova", candidates)
    assert result["result"] == "single"
    assert result["match"] == "Nərmin Əmirova"

    # Normalized exact match
    result = fuzzy_match("nermin emirova", candidates)
    assert result["result"] == "single"
    assert result["match"] == "Nərmin Əmirova"

    # Startswith normalized
    result = fuzzy_match("nermin", candidates)
    assert result["result"] == "single"
    assert result["match"] == "Nərmin Əmirova"

    # Contains normalized
    result = fuzzy_match("emirova", candidates)
    assert result["result"] == "single"
    assert result["match"] == "Nərmin Əmirova"

    # Ambiguous (two start with "nər")
    result = fuzzy_match("nər", candidates)
    assert result["result"] == "ambiguous"
    assert len(result["matches"]) >= 2

    # No match
    result = fuzzy_match("xyz123", candidates)
    assert result["result"] == "none"

    # Empty query
    result = fuzzy_match("", candidates)
    assert result["result"] == "none"


def test_resolve_user_dative_names(db_session, factories):
    """resolve_user resolves Azerbaijani dative names to users."""
    from app.ai.resolvers import resolve_user as resolver_resolve_user

    nermin = factories["create_user"](
        full_name="Nərmin Əmirova",
        email="nermin.əmirova@task-manager.local",
        role_id=3,
    )
    elmin = factories["create_user"](
        full_name="Elmin Abdullayev",
        email="elmin.abdullayev@task-manager.local",
        role_id=3,
    )

    # Create a mock current_user (admin) for the resolver
    admin = factories["create_user"](
        full_name="Admin User",
        email="admin@test.local",
        role_id=1,
        password="adminpass",
    )

    # Test dative name resolution
    result = resolver_resolve_user(
        db=db_session,
        current_user=admin,
        query="Nərminə",
    )
    assert result is not None
    assert not isinstance(result, dict)  # Not ambiguous
    assert result.full_name == "Nərmin Əmirova"
    assert result.id == nermin.id

    result = resolver_resolve_user(
        db=db_session,
        current_user=admin,
        query="Elminə",
    )
    assert result is not None
    assert not isinstance(result, dict)
    assert result.full_name == "Elmin Abdullayev"
    assert result.id == elmin.id

    # Test first-name resolution
    result = resolver_resolve_user(
        db=db_session,
        current_user=admin,
        query="Nərmin",
    )
    assert result is not None
    assert not isinstance(result, dict)
    assert result.full_name == "Nərmin Əmirova"

    # Test that non-dative query for Elmin also works
    result = resolver_resolve_user(
        db=db_session,
        current_user=admin,
        query="Elmin",
    )
    assert result is not None
    assert not isinstance(result, dict)
    assert result.full_name == "Elmin Abdullayev"


def test_resolve_project_exact_code(db_session, factories):
    """resolve_project resolves exact project code immediately."""
    from app.ai.resolvers import resolve_project as resolver_resolve_project
    from app import models as app_models

    admin = factories["create_user"](
        full_name="Admin User",
        email="admin@test.local",
        role_id=1,
        password="adminpass",
    )

    project_l026 = app_models.Project(
        code="L026",
        name="SXA L026 — İcra monitorinqi",
    )
    project_other = app_models.Project(
        code="M001",
        name="SXA M001 — Maliyyə auditi",
    )
    db_session.add_all([project_l026, project_other])
    db_session.commit()

    # Exact code match should win immediately
    result = resolver_resolve_project(
        db=db_session,
        current_user=admin,
        query="L026",
    )
    assert result is not None
    assert not isinstance(result, dict)
    assert result.code == "L026"
    assert result.name == "SXA L026 — İcra monitorinqi"

    # "SXA" is broad — both projects contain it → ambiguous
    result = resolver_resolve_project(
        db=db_session,
        current_user=admin,
        query="SXA",
    )
    assert isinstance(result, dict)
    assert result.get("ambiguous") is True
    assert len(result["projects"]) >= 2


def test_resolve_project_exact_code_case_insensitive(db_session, factories):
    """resolve_project is case-insensitive for project codes."""
    from app.ai.resolvers import resolve_project as resolver_resolve_project
    from app import models as app_models

    admin = factories["create_user"](
        full_name="Admin User",
        email="admin@test.local",
        role_id=1,
        password="adminpass",
    )

    project = app_models.Project(code="L026", name="Test Project")
    db_session.add(project)
    db_session.commit()

    result = resolver_resolve_project(
        db=db_session,
        current_user=admin,
        query="l026",
    )
    assert result is not None
    assert not isinstance(result, dict)
    assert result.code == "L026"


def test_resolve_deadline():
    """resolve_deadline parses natural-language and ISO deadlines."""
    from app.ai.resolvers import resolve_deadline
    from datetime import datetime, timezone, timedelta

    tz = timezone(timedelta(hours=4))
    now = datetime(2026, 6, 25, 10, 0, 0, tzinfo=tz)

    # "sabah" → tomorrow EOD
    result = resolve_deadline("sabah", current_datetime=now)
    assert isinstance(result, datetime)
    assert result.day == 26
    assert result.month == 6
    assert result.hour == 23

    # "bugün" → today EOD
    result = resolve_deadline("bugün", current_datetime=now)
    assert isinstance(result, datetime)
    assert result.day == 25

    # ISO date
    result = resolve_deadline("2026-07-01")
    assert isinstance(result, datetime)
    assert result.month == 7
    assert result.day == 1

    # Azerbaijani date
    result = resolve_deadline("1 iyul 2026")
    assert isinstance(result, datetime)
    assert result.day == 1
    assert result.month == 7
    assert result.year == 2026

    # "bu həftə" → Saturday EOD
    result = resolve_deadline("bu həftə", current_datetime=now)
    assert isinstance(result, datetime)

    # Unknown / unparseable
    result = resolve_deadline("bilmirəm", current_datetime=now)
    assert isinstance(result, dict)
    assert result.get("needs_more_info") is True


def test_create_task_with_resolved_fields(
    client,
    auth_headers,
    factories,
    db_session,
    app_modules,
):
    """create_task with Azerbaijani fields resolves correctly."""
    admin = factories["create_user"](
        full_name="Admin User",
        email="admin-resolver@test.local",
        role_id=1,
        password="adminpass",
    )
    assignee = factories["create_user"](
        full_name="Elmin Abdullayev",
        email="elmin@test.local",
        role_id=3,
    )

    models = app_modules["models"]
    project = models.Project(code="L026", name="SXA L026")
    db_session.add(project)
    db_session.commit()

    from app.ai.tools import create_task

    result = create_task(
        db=db_session,
        current_user=admin,
        data={
            "title": "test",
            "assignee_query": "Elminə",
            "priority": "aşağı",
            "project_query": "L026",
            "task_type": "birdəfəlik",
        },
    )

    assert result["status"] == "completed"
    assert result["tool"] == "create_task"
    task = result["task"]
    assert task["title"] == "test"
    assert task["priority"] == "low"
    assert task["task_type"] == "one_time"
    assert task["assigned_user"] == "Elmin Abdullayev"

    # Verify project assignment in DB (serialize_task doesn't include project_id)
    from uuid import UUID
    db_task = db_session.query(models.Task).filter(
        models.Task.id == UUID(task["id"])
    ).first()
    assert db_task is not None
    assert db_task.priority == "low"
    assert db_task.task_type == "one_time"
    assert db_task.project_id is not None  # project was resolved


def test_create_task_multi_turn_combined_context(
    client,
    auth_headers,
    factories,
    db_session,
    app_modules,
):
    """Multi-turn: user first gives title+assignee+priority, then project+type.

    Simulates the combined context being sent to create_task:
    Message 1: "Elminə SXA layihəsi üzrə aşağı prioritetli tapşırıq yarat: test"
    Message 2: "layihe L026 tapşırığın növü Birdəfəlik"

    Expected confirmation card: title=test, assignee=Elmin, priority=low,
    project=L026, task_type=one_time.
    """
    admin = factories["create_user"](
        full_name="Admin User",
        email="admin-multi@test.local",
        role_id=1,
        password="adminpass",
    )
    assignee = factories["create_user"](
        full_name="Elmin Abdullayev",
        email="elmin-multi@test.local",
        role_id=3,
    )

    models = app_modules["models"]
    project = models.Project(code="L026", name="SXA L026 — İcra")
    db_session.add(project)
    db_session.commit()

    from app.ai.tools import create_task

    # Simulate the AI assembling both messages into one create_task call
    result = create_task(
        db=db_session,
        current_user=admin,
        data={
            "title": "test",
            "assignee_query": "Elminə",
            "priority": "aşağı",
            "project_query": "L026",
            "task_type": "birdəfəlik",
        },
    )

    assert result["status"] == "completed"
    task = result["task"]
    assert task["title"] == "test"
    assert task["assigned_user"] == "Elmin Abdullayev"
    assert task["priority"] == "low"
    assert task["task_type"] == "one_time"

    # Verify project was resolved (check DB; serialize_task omits project_id)
    from uuid import UUID
    db_task = db_session.query(models.Task).filter(
        models.Task.id == UUID(task["id"])
    ).first()
    assert db_task is not None
    assert db_task.project_id is not None
